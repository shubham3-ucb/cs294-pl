"""Render every slide of a .pptx to PNG WITH colour emoji.
LibreOffice 7.3 PDF export silently drops Noto Color Emoji glyphs, so we split the deck
into one-slide decks and use --convert-to png (which does draw emoji). Output 1280x720.
usage: python3 slides/build/render_png.py deck.pptx outdir [prefix]
"""
import os, subprocess, sys, tempfile, shutil
from pptx import Presentation

deck, outdir = sys.argv[1], sys.argv[2]
prefix = sys.argv[3] if len(sys.argv) > 3 else os.path.splitext(os.path.basename(deck))[0]
os.makedirs(outdir, exist_ok=True)
tmp = tempfile.mkdtemp(prefix='split_')
n = len(Presentation(deck).slides)
files = []
for i in range(n):
    prs = Presentation(deck)
    lst = prs.slides._sldIdLst
    for j, sid in reversed(list(enumerate(list(lst)))):
        if j != i:
            prs.part.drop_rel(sid.rId); lst.remove(sid)
    f = os.path.join(tmp, f'{prefix}-{i+1:02d}.pptx'); prs.save(f); files.append(f)
prof = f'file:///tmp/lo_render_{os.getpid()}'
subprocess.run(['soffice', f'-env:UserInstallation={prof}', '--headless', '--convert-to', 'png',
                '--outdir', outdir] + files, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
               check=False, timeout=900)
shutil.rmtree(tmp, ignore_errors=True)
print(n, 'slides ->', outdir)
