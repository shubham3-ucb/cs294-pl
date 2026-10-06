#!/usr/bin/env python3
"""Outfit Lab deck: the record of Tuesday's class, one technical card per Git tool.

Builds slides/tuesday_app.pptx (16:9, Google Slides ready): the title, the join slide, one slide per
tool (the command, big, then WHAT IT IS, WHAT IT DOES, HOW GIT DOES IT), the paper (two slides) and the wrap.
All text comes from the app itself (app/server/steps.js: SCENES, STEPS, PAPER, WRAP_LINE;
app/public/monster.js: the start outfit), read through node, so the deck and the projector cannot drift
apart. Speaker notes are the scene's Say, Ask and Hope to hear.
Text is measured with the real fonts; the build stops if a slide would overflow.
Run:  python3 build_app_slides.py
"""
import functools
import json
import os
import re
import subprocess
from PIL import ImageFont
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

from build_simple import (FONT, MONO, INK, MUTED, PURPLE, LAVENDER, WHITE, GREEN, RED, LINE,
                          NAVY, W, H, M, OUT, new_deck, blank, text, box, chip, title_slide)

APP = os.path.join(os.path.dirname(OUT), 'app')
ORANGE = RGBColor(0xE0, 0x8A, 0x00)
AMBER = RGBColor(0xF5, 0xA5, 0x24)
PANEL = RGBColor(0xF6, 0xF6, 0xF8)
CONTENT_W = W - 2 * M
BOTTOM = H - Inches(0.45)
LABEL_SIZE = 14
WJ = '\u2060'  # word joiner: no line break here

# The static deck has no live session: name the placeholders session.js fills in.
FILL = {'{boss}': 'the boss lab', '{wallLab}': 'the first lab'}

READ_APP = """
const steps = await import('./server/steps.js');
const outfit = await import('./public/monster.js');
const pkg = JSON.parse((await import('node:fs')).readFileSync('package.json', 'utf8'));
process.stdout.write(JSON.stringify({
  tagline: pkg.description,
  start: outfit.PARTS.map((part) => outfit.PALETTE[part][outfit.START[part]]),
  scenes: steps.SCENES,
  titles: steps.STEPS.map((s) => s.title),
  paper: steps.PAPER,
  wrap: steps.WRAP_LINE,
}));
"""


def read_app():
    out = subprocess.run(['node', '--input-type=module', '-e', READ_APP], cwd=APP, check=True,
                         capture_output=True, text=True).stdout
    return json.loads(out)


def fill(s):
    for key, value in FILL.items():
        if s.startswith(key):
            s = value[:1].upper() + value[1:] + s[len(key):]
        s = s.replace(key, value)
    return s


def plain(s):
    """Drop the app's markdown (**bold**, `code`) for notes."""
    return fill(re.sub(r'\*\*|`', '', s))


def spans(s):
    """(text, font) runs: the app's `code` spans in the mono font. A word joiner after each - and /
    keeps LibreOffice from breaking inside one (`--force` would wrap as "--" and "force")."""
    parts = fill(s).replace('**', '').split('`')
    return [(part.replace('-', '-' + WJ).replace('/', '/' + WJ), MONO) if i % 2 else (part, FONT)
            for i, part in enumerate(parts) if part]


def sentences(s):
    return re.split(r'(?<=\.) ', s)


def notes(scene):
    lines = [f'Say: {plain(scene["say"])}']
    if scene.get('ask'):
        lines += [f'Ask: {plain(scene["ask"]["q"])}', f'Hope to hear: {plain(scene["ask"]["a"])}']
    return '\n\n'.join(lines)


# ---------------------------------------------------------------- measuring text
SPACING = 1.12
LEADING = 1.21 * SPACING  # Inter's line height (ascender + descender) times the paragraph spacing


@functools.lru_cache(None)
def face(family, size, bold):
    path = subprocess.run(['fc-match', '-f', '%{file}', f'{family}:{"bold" if bold else "regular"}'],
                          check=True, capture_output=True, text=True).stdout
    return ImageFont.truetype(path, size)


def width_pt(s, family, size, bold=False):
    return face(family, size, bold).getlength(s)


def line_count(s, size, width, bold=False):
    """Lines a greedy word wrap needs for s at size pt in width (EMU), measured with the real fonts."""
    room, lines, used = Emu(width).pt * 0.99, 1, 0
    for part, family in spans(s):
        space = width_pt(' ', family, size, bold)
        for word in part.split():
            w = width_pt(word, family, size, bold)
            if used and used + space + w > room:
                lines, used = lines + 1, w
            else:
                used += (space if used else 0) + w
    return lines


def text_h(s, size, width, bold=False):
    return Pt(line_count(s, size, width, bold) * size * LEADING)


# ---------------------------------------------------------------- drawing
def rich(s, x, y, w, h, line, size, bold=False):
    """One paragraph; `code` spans in the mono font."""
    p = text(s, x, y, w, h, '', size=size, spacing=SPACING).text_frame.paragraphs[0]
    p.clear()
    for part, family in spans(line):
        r = p.add_run()
        r.text = part
        r.font.name, r.font.size, r.font.bold, r.font.color.rgb = family, Pt(size), bold, INK


def stack(s, x, y, w, items, size, gap):
    """One paragraph per item, gap between them. Returns the measured height."""
    h = sum(text_h(item, size, w) for item in items) + gap * (len(items) - 1)
    tf = text(s, x, y, w, h, items, size=size, spacing=SPACING).text_frame
    for j, para in enumerate(tf.paragraphs):
        para.space_before = Pt(0) if j == 0 else gap
    return h


def rule(s, x, y, w, h, color):
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    bar.fill.solid()
    bar.fill.fore_color.rgb = color
    bar.line.fill.background()


def emoji_panel(s, x, y, w, h, emoji, size):
    box(s, x, y, w, h, fill=PANEL, radius=0.14)
    text(s, x, y, w, h, emoji, size=size, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def head(s, label, title):
    """The projector's pill and title: "STEP 2  Try two ideas at once"."""
    c = chip(s, M, Inches(0.55), label.upper())
    text(s, M + c.width + Inches(0.3), Inches(0.55), Inches(9), Inches(0.5), title, size=24,
         bold=True, anchor=MSO_ANCHOR.MIDDLE)


# ---------------------------------------------------------------- the technical card
# Like the projector's card: the command, then three rows, a label on the left. HOW GIT DOES IT sits
# on lavender. Step 1 adds the trust line under it. One type size for every card: the largest at
# which the longest card fits.
ROWS = [('is', 'WHAT IT IS'), ('does', 'WHAT IT DOES'), ('how', 'HOW GIT DOES IT')]
SIZES = [(28, 24), (26, 22), (24, 21), (22, 20), (21, 19), (20, 18)]  # (what it is / does, how)
LABEL_W, INSET, GUTTER = Inches(1.8), Inches(0.25), Inches(0.35)
TEXT_W = CONTENT_W - 2 * INSET - LABEL_W - GUTTER
ROW_PAD, HOW_PAD, HOW_GAP = Inches(0.15), Inches(0.22), Inches(0.08)
CMD_Y, CMD_H, CMD_MAX = Inches(1.25), Inches(1.0), 66
ROWS_Y = CMD_Y + CMD_H + Inches(0.25)
NOTE_SIZE, NOTE_GAP, NOTE_X = 22, Inches(0.22), Inches(0.4)


def row_style(key, sizes):
    """(type size, padding) of one row."""
    return (sizes[1], HOW_PAD) if key == 'how' else (sizes[0], ROW_PAD)


def row_h(card, key, sizes):
    size, pad = row_style(key, sizes)
    return text_h(card[key], size, TEXT_W) + 2 * pad


def note_h(note):
    return text_h(note, NOTE_SIZE, CONTENT_W - NOTE_X, bold=True)


def card_h(card, note, sizes):
    rows = sum(row_h(card, key, sizes) for key, _ in ROWS) + HOW_GAP
    return rows + (NOTE_GAP + note_h(note) if note else 0)


def card_sizes(scenes):
    cards = [(card, scene['reveal']['note']) for scene in scenes if scene['kind'] == 'reveal'
             for card in scene['reveal']['cards']]
    for sizes in SIZES:
        if all(card_h(card, note, sizes) <= BOTTOM - ROWS_Y for card, note in cards):
            return sizes
    raise SystemExit('A technical card does not fit: shorten CARDS in app/server/steps.js.')


def tool_slide(prs, app, scene, card, notes_from, sizes):
    s = blank(prs, notes=notes(notes_from))
    head(s, f'Step {scene["step"]}', app['titles'][scene['step']])
    command = card['command']
    size = min(CMD_MAX, int(Emu(CONTENT_W).pt * 97 / width_pt(command, MONO, 100, bold=True)))
    text(s, M, CMD_Y, CONTENT_W, CMD_H, command, size=size, bold=True, color=PURPLE, font=MONO,
         anchor=MSO_ANCHOR.BOTTOM, spacing=1.0)
    y = ROWS_Y
    for key, label in ROWS:
        size, pad = row_style(key, sizes)
        h = row_h(card, key, sizes)
        if key == 'how':
            y += HOW_GAP
            box(s, M, y, CONTENT_W, h, fill=LAVENDER, radius=0.12)
        else:
            rule(s, M, y, CONTENT_W, Pt(1), LINE)
        # the label sits on the first line's baseline
        drop = Pt((size - LABEL_SIZE) * 0.97)
        text(s, M + INSET, y + pad + drop, LABEL_W, Inches(0.3), label, size=LABEL_SIZE, bold=True,
             color=PURPLE if key == 'how' else MUTED)
        rich(s, M + INSET + LABEL_W + GUTTER, y + pad, TEXT_W, h - 2 * pad, card[key], size)
        y += h
    note = scene['reveal']['note']
    if note:
        y += NOTE_GAP
        rule(s, M, y, Inches(0.08), note_h(note), AMBER)
        rich(s, M + NOTE_X, y, CONTENT_W - NOTE_X, note_h(note), note, NOTE_SIZE, bold=True)


# ---------------------------------------------------------------- slides
def cover(prs, app):
    s = title_slide(prs, 'Outfit Lab', sentences(app['tagline']),
                    'CS294 · Git Week · Tuesday · Shubham & Ananya',
                    "In class the projector page runs the lesson, and the console's Next button moves "
                    'through it. This deck is the record: one technical card per Git tool, the same text '
                    'as the app.')
    # the start outfit, as the app draws it: hat, glasses, top, shoes
    pw, ph, gap = Inches(1.9), Inches(0.95), Inches(0.12)
    y0 = (H - Inches(0.6) - 4 * ph - 3 * gap) / 2
    for i, e in enumerate(app['start']):
        emoji_panel(s, W - M - pw, y0 + i * (ph + gap), pw, ph, e, 44)


def join(prs, scene):
    s = blank(prs, notes=notes(scene))
    q = Inches(4.4)
    qx, qy = W - M - q, (H - q) / 2
    box(s, qx, qy, q, q, fill=WHITE, line=LINE, dash=True, radius=0.06)
    text(s, qx, qy, q, q, 'QR', size=32, color=MUTED, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE)
    left = qx - M - Inches(0.4)
    text(s, M, Inches(1.9), left, Inches(0.5), 'Outfit Lab', size=22, bold=True, color=PURPLE)
    text(s, M, Inches(2.5), left, Inches(0.9), 'Open: ________', size=48, bold=True)
    text(s, M, Inches(3.7), left, Inches(1.6), sentences(scene['line']), size=28, color=MUTED,
         spacing=1.25)


def reveal(prs, app, scene, sizes, task):
    """One slide per card. Step 4 has two: the first takes the task's notes (the Wall, clone, push)."""
    cards = scene['reveal']['cards']
    for i, card in enumerate(cards):
        tool_slide(prs, app, scene, card, scene if i == len(cards) - 1 else task, sizes)


def paper(prs, app, scene):
    """Two slides. First: Good / Bad / Ugly, then the one message (the paper's claim, for analysts) and its
    trade-off for developers. Second: what the class lived, and the pairs question. Each band starts below the last."""
    s = blank(prs, notes=notes(scene))
    p = app['paper']
    text(s, M, Inches(0.45), CONTENT_W, Inches(0.4), f"Tuesday's paper · {p['source']}", size=18,
         color=MUTED)
    text(s, M, Inches(0.8), CONTENT_W, Inches(0.6), p['title'], size=32, bold=True)
    gap = Inches(0.4)
    cw = (CONTENT_W - 2 * gap) / 3
    y = Inches(2.15)
    col_h = 0
    for i, (label, color) in enumerate([('GOOD', GREEN), ('BAD', ORANGE), ('UGLY', RED)]):
        x = M + i * (cw + gap)
        rule(s, x, Inches(1.6), cw, Inches(0.06), color)
        text(s, x, Inches(1.75), cw, Inches(0.3), label, size=18, bold=True, color=color)
        col_h = max(col_h, stack(s, x, y, cw, p[label.lower()], 18, Pt(8)))
    y += col_h + Inches(0.35)
    message = text_h(p['message'], 36, CONTENT_W, bold=True)
    tradeoff = text_h(p['tradeoff'], 20, CONTENT_W)
    if y + message + Inches(0.1) + tradeoff > BOTTOM:
        raise SystemExit('The paper slide overflows: shorten PAPER in app/server/steps.js.')
    text(s, M, y, CONTENT_W, message, p['message'], size=36, bold=True, color=PURPLE)
    y += message + Inches(0.1)
    text(s, M, y, CONTENT_W, tradeoff, p['tradeoff'], size=20, spacing=SPACING)

    # What they lived, linked to the paper; then the question pairs answer in the app.
    s = blank(prs, notes=notes(scene))
    head(s, 'The paper', 'What you lived')
    y = Inches(1.5)
    y += stack(s, M, y, CONTENT_W, [plain(l['text']) for l in p['lived']], 22, Pt(10)) + Inches(0.5)
    text(s, M, y, CONTENT_W, Inches(0.3), 'PAIRS, IN THE APP', size=LABEL_SIZE, bold=True, color=MUTED)
    y += Inches(0.35)
    question = text_h(plain(scene['ask']['q']), 24, CONTENT_W, bold=True)
    if y + question > BOTTOM:
        raise SystemExit('The paper slide overflows: shorten PAPER in app/server/steps.js.')
    text(s, M, y, CONTENT_W, question, plain(scene['ask']['q']), size=24, bold=True, color=PURPLE, spacing=SPACING)


def wrap(prs, app, scene):
    s = blank(prs, NAVY, notes(scene))
    for x, y, d in [(1.2, 0.8, .05), (3.4, 1.6, .04), (5.9, .6, .06), (8.3, 1.3, .04),
                    (10.6, .7, .05), (12.1, 1.9, .04), (2.2, 6.4, .04), (11.4, 6.1, .05),
                    (7.2, 6.7, .04), (0.7, 3.9, .03), (12.6, 4.2, .03)]:
        star = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
        star.fill.solid()
        star.fill.fore_color.rgb = WHITE
        star.line.fill.background()
    text(s, M, Inches(2.1), CONTENT_W, Inches(3.3), sentences(app['wrap']), size=60, bold=True,
         color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, spacing=1.2)


def flatten(prs):
    """Drop each shape's theme style: LibreOffice draws its shadow despite an empty effectLst."""
    for slide in prs.slides:
        for shp in slide.shapes:
            style = shp._element.find(qn('p:style'))
            if style is not None:
                shp._element.remove(style)


def build():
    """Title, join, a slide per tool card (the reveals, Steps 1-6), the paper, the wrap."""
    app = read_app()
    scenes = app['scenes']
    sizes = card_sizes(scenes)
    tasks = {scene['step']: scene for scene in scenes if scene['kind'] == 'task'}
    prs = new_deck()
    cover(prs, app)
    for scene in scenes:
        kind = scene['kind']
        if kind == 'join':
            join(prs, scene)
        elif kind == 'reveal':
            reveal(prs, app, scene, sizes, tasks[scene['step']])
        elif kind == 'paper':
            paper(prs, app, scene)
        elif kind == 'wrap':
            wrap(prs, app, scene)
    flatten(prs)
    return prs, sizes


if __name__ == '__main__':
    path = os.path.join(OUT, 'tuesday_app.pptx')
    prs, sizes = build()
    prs.save(path)
    print('saved', path, f'({len(prs.slides)} slides, card type {sizes[0]}/{sizes[1]} pt)')
