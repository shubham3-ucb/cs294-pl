#!/usr/bin/env python3
"""Shared minimal slide helpers for the deck builders (build_app_slides.py, build_thursday.py):
one idea per slide, big type, lots of white space."""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.dirname(HERE)

FONT, MONO = 'Inter', 'Noto Sans Mono'
INK = RGBColor(0x11, 0x11, 0x11)
MUTED = RGBColor(0x70, 0x70, 0x78)
PURPLE = RGBColor(0x8B, 0x3D, 0xFF)
LAVENDER = RGBColor(0xF4, 0xEC, 0xFB)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREEN = RGBColor(0x16, 0xA3, 0x4A)
RED = RGBColor(0xDC, 0x26, 0x26)
LINE = RGBColor(0xD9, 0xD9, 0xDE)
NAVY = RGBColor(0x0B, 0x16, 0x30)
YELLOW = RGBColor(0xFF, 0xE0, 0x66)

W, H = Inches(13.333), Inches(7.5)
M = Inches(0.9)  # side margin


def new_deck():
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    return prs


def blank(prs, bg=WHITE, notes=''):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = bg
    if notes:
        s.notes_slide.notes_text_frame.text = notes
    return s


def text(s, x, y, w, h, lines, size=32, color=INK, bold=False, font=FONT,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.1):
    """lines: str or list of str | (str, dict-of-overrides)."""
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.vertical_anchor = anchor
    for side in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'):
        setattr(tf, side, 0)
    if isinstance(lines, str):
        lines = [lines]
    for i, ln in enumerate(lines):
        opts = {}
        if isinstance(ln, tuple):
            ln, opts = ln
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        r = p.add_run()
        r.text = ln
        f = r.font
        f.name = opts.get('font', font)
        f.size = Pt(opts.get('size', size))
        f.bold = opts.get('bold', bold)
        f.color.rgb = opts.get('color', color)
    return tb


def box(s, x, y, w, h, fill=WHITE, line=None, radius=0.12, dash=False, lw=1.5):
    shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    shp.adjustments[0] = radius
    shp.shadow.inherit = False
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(lw)
        if dash:
            from pptx.enum.dml import MSO_LINE_DASH_STYLE
            shp.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    shp.text_frame.text = ''
    return shp


def chip(s, x, y, label, fill=INK, color=WHITE, size=18, w=None):
    w = w or Inches(0.45 + 0.17 * len(label))
    c = box(s, x, y, w, Inches(0.5), fill=fill, radius=0.5)
    tf = c.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    r.font.name, r.font.size, r.font.bold = FONT, Pt(size), True
    r.font.color.rgb = color
    return c


# ---------------------------------------------------------------- layouts
def title_slide(prs, title, sub, foot, notes):
    s = blank(prs, notes=notes)
    text(s, M, Inches(2.2), W - 2 * M, Inches(1.4), title, size=64, bold=True,
         anchor=MSO_ANCHOR.BOTTOM)
    text(s, M, Inches(3.85), W - 2 * M, Inches(1.2), sub, size=26 if len(sub) > 40 else 32,
         color=MUTED)
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, H - Inches(0.6), W, Inches(0.6))
    bar.fill.solid(); bar.fill.fore_color.rgb = PURPLE; bar.line.fill.background()
    text(s, M, H - Inches(0.6), W - 2 * M, Inches(0.6), foot, size=18, color=WHITE,
         anchor=MSO_ANCHOR.MIDDLE)
    return s
