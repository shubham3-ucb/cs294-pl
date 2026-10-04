"""Builds slides/tuesday_monster_lab.pptx (Tuesday 10/13, "Monster Lab") from design/spec.md
section 5 (T1-T34; B1 is a normal slide right after T14; B2-B3 are hidden backups at the end),
using the shared style library. style.py is not modified: the Tuesday-only overrides
(card() with a 13 pt ID-tag floor, a pause bar with the 💬 in its own box, talk() with an
optional meaning size) live in this file and are patched into the style module so that
fork(), card_stack(), conflict(), terminal() etc. use them too. Then renders PNGs for review.

    python3 build_tuesday.py            # build + render
    python3 build_tuesday.py --no-render
"""
import os
import shutil
import sys

sys.path.insert(0, '/home/shagarw_google_com/shubham/git_course/slides/build')
import style as _style                  # noqa: E402
from style import *                     # noqa: E402,F401,F403
from style import _set_notes, _zone, _emo, _is_emoji, _PANEL_MARK     # noqa: E402

ROOT = '/home/shagarw_google_com/shubham/git_course'
OUT = f'{ROOT}/slides/tuesday_monster_lab.pptx'
REN = f'{ROOT}/slides/renders/tuesday'
REVIEW = '/tmp/tuesday_review.pptx'      # same deck with hidden slides un-hidden (for PNGs)

# BEFORE CLASS: set the real Mentimeter access code of the "Git Tue · Monster Lab" presentation
# and save its QR PNG at MENTI_QR, then rebuild. T2, T14 and T31 pick both up automatically.
# Until then the slides show a dashed 'QR' box and '____' (or paste them in Google Slides).
MENTI_CODE = '____'                              # e.g. '1234 5678'
MENTI_QR = f'{ROOT}/slides/assets/menti_qr_tue.png'

prs = new_deck()
BACKUPS = []

# official card set (spec 3.3) and later toy IDs (spec 3.4 + cheat sheet)
C1 = dict(cid='#1', panels=('smiley', 'box', 'sticks'), parent='—')
C2 = dict(cid='#2', panels=('cat', 'robot', 'sticks'), parent='#1')
C3 = dict(cid='#3', panels=('smiley', 'superhero', 'tentacles'), parent='#1')
C4 = dict(cid='#4', panels=('cat', 'cape', 'tentacles'), parent='#2 + #3')
FOUR = [C1, C2, C3, C4]
DRT3 = dict(cid='drt3', panels=('dragon', 'cape', 'tentacles'), parent='#4')
CRR3 = dict(cid='crr3', panels=('cat', 'cape', 'skates'), parent='#4')


# =============================================================== Tuesday overrides of style.py
TAG_PT = 13            # card ID tags never go below 13 pt Mono Bold (spec: card metadata)
CF_PT = 11             # came-from line floor
TAG_H = 0.22           # min ID-tag height (fits a 13 pt Mono line)
PB_EMO = 0.60          # pause bar: the question starts this far right of the 💬


def card(sl, x, y, cid='#1', panels=('smiley', 'box', 'sticks'), parent=None, author=None,
         s=1.0, state='committed', marks=None):
    """style.card() with legible metadata at every scale: the ID tag is >= 13 pt Mono Bold
    (box widened to fit) and the came-from line >= 11 pt. On small cards the three panels
    shrink to make room; at s >= 0.85 the geometry is identical to style.card()."""
    marks = marks or {}
    w, h = CARD_W * s, CARD_H * s
    pad = 0.12 * s
    ghost = state == 'ghost'
    meta = MUTED if ghost else INK2          # ghost metadata: faded, but still readable
    body_line = {'committed': CARD_LINE, 'staged': CARD_LINE, 'sketch': SKETCH_LINE,
                 'ghost': GHOST}[state]
    box(sl, x, y, w, h, fill=(GHOST_T if ghost else WHITE), line=body_line, lw=1.5,
        radius=0.12 * s, dash=state in ('sketch', 'ghost'))
    # --- ID tag (font does NOT scale below TAG_PT)
    tsize = max(TAG_PT, round(14 * s, 1))
    th = max(0.34 * s, TAG_H)
    tw = min(w - 2 * pad, max(0.50 * s, text_width(cid or 'drt3', tsize, True, MONO)
                              + 0.06 + 0.20 * s))
    if state in ('committed', 'ghost'):
        tg = chip(sl, x + pad, y + pad, tw, th, cid, fill=(None if ghost else INK),
                  line=(GHOST if ghost else None), lw=1, color=(MUTED if ghost else WHITE),
                  size=tsize, font=MONO, radius=0.06 * s)
        tf = tg.text_frame
        tf.margin_left = tf.margin_right = Inches(0.02)
        tf.margin_top = tf.margin_bottom = Inches(0)
    else:
        tw = 0.70 * s
        box(sl, x + pad, y + pad, tw, th, line=SKETCH_LINE, lw=1, radius=0.06 * s, dash=True)
    # --- author, right of the tag
    if author:
        aw = w - 2 * pad - tw - 0.05 * s
        asz = fit_size(author, aw, th, [13, 12, 11], max_lines=1, warn=False)
        text(sl, x + pad + tw + 0.05 * s, y + pad, aw, th, author, size=asz,
             color=meta, align='r', anchor='m', markup=False)
    # --- vertical layout: tag | 3 panels | came-from row (same for every card of scale s)
    chh = max(0.26 * s, CF_PT * 1.25 / 72)
    cf_cy = min(y + 2.25 * s, y + h - chh / 2 - 0.012)
    py0 = y + max(0.56 * s, pad + th + 0.025)
    k = (min(y + 2.12 * s, cf_cy - chh / 2) - py0) / (1.56 * s)
    ph, pitch = 0.48 * s * k, 0.54 * s * k
    for i in range(3):
        emo = _emo(panels[i]) if panels and i < len(panels) else ''
        mk = marks.get(i)
        fill, ln, lw, dash = PANEL, PANEL_LINE, 0.75, False
        if mk in _PANEL_MARK:
            fill, ln, lw, dash = _PANEL_MARK[mk]
            if mk == 'conflict' and not emo:
                emo = '❓'
        if ghost:
            fill, ln, lw, dash, emo = GHOST_T, PANEL_LINE, 0.75, False, ''
        p = box(sl, x + pad, py0 + i * pitch, w - 2 * pad, ph, fill=fill, line=ln, lw=lw,
                dash=dash, radius=0.05 * s)
        if emo:
            esz = (26 if len([c for c in emo if _is_emoji(c)]) < 2 else 22) * s * k
            label(p, emo, size=round(esz, 1), bold=False, markup=False)
    # --- came-from line (>= CF_PT; a two-parent line drops its spaces before shrinking)
    if parent is not None:
        def runs(par):
            return [['—']] if par in ('—', '-', '') else [[('← ', {'font': FONT}),
                                                           (par, {})]]
        cw_ = w - 0.08                      # centred, so it may use the card's full width
        sizes = [max(CF_PT, round(13 * s * kk)) for kk in (1, .9, .8, .7)]
        cf = runs(parent)
        if lines_needed(cf, cw_, sizes[-1], font=MONO) > 1 and ' + ' in parent:
            cf = runs(parent.replace(' + ', '+'))
        csz = fit_size(cf, cw_, 1.0, sizes, font=MONO, max_lines=1, warn=False)
        text(sl, x + 0.04, cf_cy - chh / 2, cw_, chh, cf, size=csz, color=meta, font=MONO,
             align='c', anchor='m', markup=False)
    return dict(x=x, y=y, w=w, h=h, l=x, r=x + w, t=y, b=y + h, cx=x + w / 2, cy=y + h / 2,
                row_y=[py0 + i * pitch + ph / 2 for i in range(3)], s=s, id=cid)


def pause_bar(sl, question, y_bottom=6.92, x=M, w=CW):
    """Lavender pause bar. The 💬 sits in its own box and the question in a second box
    PB_EMO to its right, so a two-line question keeps one clean left edge (hanging indent).
    24 pt on one line; else 22 pt on <= 2 lines. Returns the bar's top y."""
    inner = w - 0.60 - PB_EMO
    size, n = 24, lines_needed(question, inner, 24)
    if n > 1:
        size = 22
        n = lines_needed(question, inner, 22)
        if n > 2:
            size = fit_size(question, inner, 0.80, (20, 18), max_lines=2)
            n = 2
    h = 0.72 if n == 1 else 1.02
    top = y_bottom - h
    box(sl, x, top, w, h, fill=LAVENDER, radius=0.14)
    text(sl, x + 0.30, top, PB_EMO, h, '💬', size=size, anchor='m', markup=False)
    text(sl, x + 0.30 + PB_EMO, top, inner, h, question, size=size, anchor='m', ls=1.05,
         balance=('widow' if n > 1 else False))
    return top


def talk(prs, command, meaning, kicker='You just invented', pause=None, step=None,
         total=9, notes=None, msize=None):
    """style.talk() with an optional fixed meaning size (msize)."""
    sl = blank(prs)
    text(sl, M, TOP, CW, 0.50, kicker, size=24, color=INK2, anchor='m')
    csz = fit_size(command, CW, 1.05, (54, 48, 44), bold=True, font=MONO, max_lines=1,
                   markup=False)
    text(sl, M, 1.12, CW, 1.05, command, size=csz, bold=True, color=PURPLE, font=MONO,
         anchor='m', markup=False)
    msz = msize or fit_size(meaning, CW, 1.20, (32, 30, 28), ls=1.12)
    text(sl, M, 2.30, CW, 1.20, meaning, size=msz, ls=1.12, balance='widow')
    mlines = lines_needed(balanced(meaning, CW, msz, mode='widow'), CW, msz)
    sl.meaning_b = 2.30 + mlines * msz * 1.21 * 1.12 / 72
    zt = sl.meaning_b + 0.25
    sl.pause_t = pause_bar(sl, pause) if pause else 7.10
    sl.zone = _zone(M, zt, R, sl.pause_t - 0.20 if pause else 6.90)
    if isinstance(step, int):
        progress(sl, step, total, y=7.12)
    return _set_notes(sl, notes)


# style's layouts (fork, card_stack, conflict, terminal, ...) look these up at call time
_style.card = card
_style.pause_bar = pause_bar
_style.talk = talk


# =============================================================== local helpers
def band_scene(sl, top, bottom, zone_w=2.6, wall_w=2.6, left='Lab 1 table',
               right='Lab 2 table'):
    """Two tables + the Wall inside a TALK slide's diagram band, with wide (2.0 in) gaps so
    arrows and their labels have room. Returns dict(left, wall, right, gap_l, gap_r)."""
    gap = (CW - 2 * zone_w - wall_w) / 2
    h = bottom - top
    lz = table_zone(sl, M, top, zone_w, h, left)
    wz = wall(sl, M + zone_w + gap, top, wall_w, h)
    rz = table_zone(sl, R - zone_w, top, zone_w, h, right)
    return dict(left=lz, wall=wz, right=rz, gap_l=(lz['r'], wz['l']), gap_r=(wz['r'], rz['l']))


S_BAND = 0.43          # card scale inside the band scenes (T18, T23)
STACK_STEP = 0.25      # absolute cascade step: keeps every 13 pt ID tag in a stack readable
STACK_DY = STACK_STEP / S_BAND
BAND_MEANING = 28      # T18/T23: 28 pt two-line meaning leaves air above the wall scene


def band_bounds(sl):
    """T18/T23 wall-scene band: 0.20 in below the meaning's zone top (pill clears the text
    by ~0.25 in), 0.12 in above the pause bar."""
    return sl.zone['t'] + 0.20, sl.pause_t - 0.12


def stack_size(n=4, s=S_BAND):
    return CARD_W * s + 0.10 * (n - 1), CARD_H * s + STACK_DY * s * (n - 1)


def underbrace(sl, g, color=INK2, drop=0.14, tick=0.10):
    """Square bracket under a card: 'everything on this card'. Returns its bottom y."""
    y = g['b'] + drop
    line(sl, g['l'], y - tick, g['l'], y, color=color, w=1.5)
    line(sl, g['l'], y, g['r'], y, color=color, w=1.5)
    line(sl, g['r'], y - tick, g['r'], y, color=color, w=1.5)
    line(sl, g['cx'], y, g['cx'], y + tick, color=color, w=1.5)
    return y + tick


def rows(sl, x, y, w, items, marks=None, size=28, mark_w=None, gap=0.10, ls=1.08):
    """Numbered rows at a FIXED pitch (style.numbered() grows a row that contains a mono
    run, which makes the gaps uneven). Items must fit on one line each."""
    mark_w = mark_w or size / 72 * 1.2
    lh = size * 1.21 * ls / 72
    for i, it in enumerate(items):
        yy = y + i * (lh + gap)
        mk = marks[i] if marks else str(i + 1)
        text(sl, x, yy, mark_w, lh + 0.02, mk, size=size, bold=True, color=PURPLE, ls=ls,
             markup=False)
        text(sl, x + mark_w, yy, w - mark_w, lh + 0.10, it, size=size, ls=ls)
    return y + len(items) * (lh + gap)


def task_slide(prs, step, instruction, lines, minutes, notes):
    """style.task() with its sub-steps drawn by rows() (same geometry, even spacing)."""
    sl = task(prs, step, instruction, (), minutes=minutes, notes=notes)
    rows(sl, M, 4.00, CW, list(lines))
    return sl


def hide(sl):
    sl._element.set('show', '0')
    BACKUPS.append(sl)
    return sl


# =============================================================== T1-T2
title_slide(prs, 'Monster Lab', 'Build Git out of paper',
            topic='Git Week · Implementation Day',
            notes=[
                "No Git lecture today. You'll rebuild Git out of paper, and every Git tool "
                "shows up only after you've felt the problem it fixes.",
                "Sit with your lab (3 labs of 3-4: Lab 1, Lab 2, Lab 3). Phones out for Menti.",
            ])

menti(prs, 'You type `git branch experiment`. What does Git create?',
      ['A copy of all your files', 'A copy of all your commits', 'One new commit',
       'One tiny file holding one ID'], code=MENTI_CODE, qr=MENTI_QR,
      notes=[
          "Open this poll as students arrive. \"Gut answer. We won't reveal it; you'll "
          "discover it with paper.\"",
          "Hide results. The answer, D (one tiny file holding one ID), is revealed on T10 "
          "(git branch). Mark D as correct in Menti.",
      ])

# =============================================================== PART 1 · One table
divider(prs, 'One table', PINK, kicker='PART 1',
        notes='(5 seconds) "Part 1: one team, one table."')

# T4 · Step 0
task_slide(prs, 0, 'Draw one monster. Together. One sheet.',
     ['Open your secret slip', 'Pencils only', 'Go!'], minutes='1:30',
     notes=[
         'Hand out the A3 sheets, pencils, and Step-0 slips (P1-P4). Start the phone timer: '
         '1:30. Call "Forty-five seconds!" at 0:45 (that is P3\'s cue to erase the horns), '
         'then stop at 1:30.',
         'PAUSE / debrief (~90 s): "What did the face look like at 0:30? Who added the '
         'horns? Get me the 0:30 monster back." Nobody can.',
         'Then: "Invent ONE rule that would make these answerable."',
         'Nudge: "What if erasing were illegal?"',
         'Check: "Does a card store the change, or the whole monster?" (the whole monster, a '
         'snapshot; unchanged parts can be shared)',
     ])

# T5 · TALK git commit (card states)
sl = talk(prs, 'git commit', 'Never erase. Every change = a new card of the **whole** monster.',
          pause='You pencilled a new FACE and a new BODY. Only FACE is ready. What goes on '
                'the next card?', step=0,
          notes=[
              '"A commit is a snapshot, not a diff. Pencil = working copy, ink = staged, '
              'writing the ID = committed."',
              'Pause answer: ink only FACE. The next card gets the new FACE and the OLD BODY, '
              'because a card is always the whole monster. That is `git add -p` (or `git add '
              'face.txt`). The index is the next card, being drafted.',
              '`git diff` = pencil vs ink. `git diff --staged` = ink vs the last card.',
              'If a Model Critic says "but packfiles store diffs!": "Write it down. You\'re '
              'right, and you get the floor at the end."',
              'Each lab names its Model Critic (whoever already uses Git daily). Hand out '
              'cards, markers, critic sheets: "Markers only from now on. One word + a '
              '5-second doodle per panel."',
          ])
z = sl.zone
s = round(min(0.85, (z['h'] - 0.06) / CARD_H), 3)
y0 = z['t'] + (z['h'] - CARD_H * s) / 2
cw = CARD_W * s
units = [('sketch', 'working copy', '', None), ('staged', 'git add', '', None),
         ('committed', 'git commit', '#1', '—')]
lws = [text_width(u[1], 18, True, MONO) + 0.05 for u in units]
ugap, lgap = 0.70, 0.20
total = sum(cw + lgap + lw for lw in lws) + ugap * (len(units) - 1)
xx = (W - total) / 2
for (st, lab, cid, par), lw in zip(units, lws):
    g = card(sl, xx, y0, cid, ('smiley', 'box', 'sticks'), parent=par, s=s, state=st)
    text(sl, g['r'] + lgap, g['cy'] - 0.22, lw + 0.1, 0.44, lab, size=18, bold=True,
         color=PURPLE, font=MONO, anchor='m', markup=False)
    xx = g['r'] + lgap + lw + ugap

# T6 · goals
content(prs, "Today's goals", lead='After today, we should be confident you could:', items=[
    'Explain a commit, a branch, and HEAD', 'Predict what conflicts in a merge',
    'Fix a rejected push', 'Choose revert, reset, or rebase, and say what each destroys'],
    notes=[
        '"The exit ticket checks all four. You can\'t pass it by memorising commands, only '
        'with the model."',
        '(1 minute. Do not explain any of the terms yet.)',
    ])

# T7 · Step 1
task_slide(prs, 1, "Another lab's history. Put it in order.",
     ['Which card is newest?', 'Horns added, or removed?', 'Fix it: what must every card say?'],
     minutes='3:30',
     notes=[
         'Envelopes out (W, X, Y, Z), fronts up. After about 90 seconds they realise it '
         'can\'t be decided: both W-X-Y-Z and Z-Y-X-W fit the fronts.',
         'The invention we want: author + "came from __". Nudge: "What ONE line on each card '
         'would survive shuffling?" Only then: "Flip them. Draw the arrows."',
         'Expected: horns added (W -> X); Y and Z are both tips, so there\'s no single newest. '
         '"Two lines of work: that\'s Step 2."',
         'Then each lab draws #1: smiley / box / sticks, with "#1 · by __ · came from: —".',
         'Wrong answer, timestamps: "Laptop clocks disagree. And does a time tell you what it '
         'was built FROM?"',
     ])

# T8 · TALK git log
sl = talk(prs, 'git log', 'Each card points to its parent. History = follow the arrows back.',
          pause='Why do the arrows point backward, never forward?', step=1,
          notes=[
              '"Git keeps no separate history list. `log` starts at a tip and walks the '
              'parent pointers."',
              'Pause answer: a card never changes, so an old card can\'t learn about its '
              'future children. Pointers can only go from child to parent.',
              'The first commit has no parent: that\'s the root.',
              'Card Z also points to W, so there are two tips. Next step.',
          ])
s = 0.80
z = sl.zone
cy0 = card_top(z, s)
xs0 = W / 2 - 2.85 - CARD_W * s / 2
gs = []
for i, (cid, pan, au, par) in enumerate([('W', ('smiley', 'box', 'sticks'), 'Ken', '—'),
                                         ('X', ('horns', 'box', 'sticks'), 'Barbara', 'W'),
                                         ('Y', ('horns', 'box', 'wheels'), 'Margaret', 'X')]):
    gs.append(card(sl, xs0 + i * 2.85, cy0, cid, pan, parent=par, author=au, s=s))
for c1, c0 in zip(gs[1:], gs[:-1]):
    link(sl, c1, c0)
row_labels(sl, gs[0], size=13)

# T9 · Step 2
task_slide(prs, 2, 'Two risky ideas. Keep the good monster safe.',
     ['Split into two pairs', 'Open your idea card. Draw ONE card.',
      'Anyone finds the approved monster in 3 seconds'], minutes='4:00',
     notes=[
         'Idea cards: CAT-ROBOT (card #2: FACE cat, BODY robot) and SUPERHERO (card #3: BODY '
         'superhero, LEGS tentacles). Both start from #1.',
         'At 2:30 Ananya plays the Client: "Client\'s here! Show me the approved monster AND '
         'each idea\'s latest card. Now!"',
         'The invention we want: sticky flags (`main` on #1, `cat-robot`, `superhero` on each '
         'idea\'s latest card).',
         'Nudge: "Which card is approved? How would the Client know without asking you?"',
         'Check: "If the cat-robot pair draws one more card, what\'s its came-from, and which '
         'flags move?" (the cat-robot card; only `cat-robot` moves, `main` stays)',
     ])

# T10 · TALK git branch (fork)
sl = talk(prs, 'git branch', 'A branch is a flag on a card. HEAD = you are here.',
          pause="Pin HEAD on a card with no flag, draw a card, then switch back. Where's your "
                "new card?", step=2,
          notes=[
              'LIVE, 60 s, inside .git (repo monster/):',
              '  cat .git/HEAD            -> ref: refs/heads/cat-robot',
              '  git rev-parse HEAD > .git/refs/heads/hack && git branch   -> a new branch '
              '"hack", made without any branch command. wc -c .git/refs/heads/hack -> 41 '
              'bytes.',
              '"HEAD decides your next commit\'s parent. Committing moves the flag HEAD points '
              'to."',
              'Pause answer: a pin on a flagless card = DETACHED HEAD. The new card has no '
              'flag, so nothing leads to it. Git warns you (verified): "Warning: you are '
              'leaving 1 commit behind, not connected to any of your branches:". `git switch '
              '-c wheels` would have saved it. Otherwise only the reflog remembers it (Step 7).',
              'Optional demo: git switch --detach main; edit; git commit -am wheels; git '
              'switch cat-robot. Afterwards: git branch -D hack.',
              'SAY: "New idea = new flag. Zero copied cards." Then: "So the poll answer is '
              'D." Show the T2 Menti results now.',
          ])
# fork centred between the meaning and the (two-line) pause bar, and across the slide
s, fgap, fdx = 0.48, 0.12, 2.95
fh_ = 2 * CARD_H * s + fgap
ftop = sl.meaning_b + (sl.pause_t - sl.meaning_b - fh_) / 2
f1w = max(1.20 * s, chip_w('cat-robot', 14, True, pad=0.12))
grp_w = fdx + CARD_W * s - 0.15 + f1w + 0.06 + 0.13 + 0.86      # root card .. HEAD pin
g0, (t1, t2) = fork(sl, C1, [C2, C3], x=(W - grp_w) / 2, top=ftop, s=s, dx=fdx, gap=fgap)
flag(sl, g0, 'main', size=14)
f1 = flag(sl, t1, 'cat-robot', edge='right', size=14)
flag(sl, t2, 'superhero', edge='right', size=14)
head_pin_on(sl, f1, side='right')

# T11 · Step 3
task_slide(prs, 3, 'The Client wants ONE monster with both ideas.',
     ['One new card. Never erase.', 'Which panels are easy? Which need a talk?',
      'Disagree? Red ❓, then decide together.'], minutes='5:00',
     notes=[
         'The invention we want: a card with TWO came-froms, plus a panel-by-panel '
         'comparison with #1 (the card both pairs started from).',
         'Expected result: FACE cat, LEGS tentacles, BODY ❓ -> robot with a cape.',
         'Nudge: "Who changed each panel since the card you both started from?"',
         'Wrong answers: "Take the newest card" -> "Then whose work did you just throw '
         'away?"  "Everything conflicts" -> "Did superhero touch FACE?"',
         'Check: "Cover #1. Can you still prove FACE isn\'t a conflict?" (No: cat vs smiley '
         'looks like a disagreement. A 3-way merge needs the ancestor.)',
         'If a lab is stuck: put the rescue set\'s #1, #2 and #3 side by side.',
     ])

# T12 · CONFLICT git merge (the key visual)
conflict(prs, pause='Why does Git need BASE? Why not just compare the two?', step=3,
         notes=[
             '"Three-way merge: for each part, Git asks who changed it since the shared '
             'ancestor (the merge base)."',
             'Pause answer: without BASE, cat vs smiley looks like a disagreement, and every '
             'panel would conflict. With BASE: only one side touched FACE -> take it; only one '
             'side touched LEGS -> take it; both touched BODY differently -> you decide.',
             '"We resolved BODY by hand (robot with a cape), on a card with two parents."',
             'Keep the word "differently": exit question Q-a depends on it (identical changes '
             'on both sides merge cleanly).',
         ])

# T13 · TERMINAL real git merge
sl, gl, gr = terminal_split(prs, 'Same merge, real Git', [
    ('cmd', 'git switch cat-robot'),
    ('cmd', 'git merge superhero'),
    'Auto-merging monster.txt',
    ('err', 'CONFLICT (content): Merge'),
    ('err', '  conflict in monster.txt'),
], [
    'face: cat', '', ('mark', '<<<<<<< HEAD'), 'body: robot', ('mark', '======='),
    'body: superhero', ('mark', '>>>>>>> superhero'), '', 'legs: tentacles',
], right_title='monster.txt', widths=(5.60, 5.90),
    notes=[
        'Run it live (Appendix A, repo monster/). "In this file each panel is one line, with '
        'a blank line between panels."',
        '  git log --oneline --graph --all',
        '  git merge superhero      -> body-only CONFLICT;  cat monster.txt',
        'Resolve:  printf \'face: cat\\n\\nbody: robot with a cape\\n\\nlegs: tentacles\\n\' '
        '> monster.txt && git add monster.txt && git commit -m "merge: robot with a cape"',
        'Then git log --oneline --graph: the diamond.',
        'git cat-file -p HEAD shows TWO `parent` lines: "exactly your paper card".',
        'This slide is also the backup if the live demo fails.',
    ])
term_chip(sl, gr, 0, 11)
term_chip(sl, gr, 8, 17)

# T14 · MENTI predict
menti(prs, 'I delete the blank lines and merge again. What conflicts?',
      ['BODY only', 'FACE + BODY', 'All three lines', 'Nothing'], code=MENTI_CODE, qr=MENTI_QR,
      notes=[
          'Vote (hide results). The next slide shows the answer (B1); don\'t run it live.',
          'Answer C: Git merges LINES. Changed lines that touch each other form one hunk, so '
          'the three edits become one conflict.',
          'False alarm: nobody actually disagreed about FACE or LEGS.',
      ])

# B1 · TERMINAL answer to T14 (a normal slide, so present mode reaches it: Google Slides
# turns hidden slides into skipped slides)
terminal(prs, [
    ('cmd', 'git merge superhero'),
    ('err', 'CONFLICT (content): Merge conflict in monster.txt'),
    ('mark', '<<<<<<< HEAD'),
    'face: cat',
    'body: robot',
    'legs: sticks',
    ('mark', '======='),
    'face: smiley',
    'body: superhero',
    'legs: tentacles',
    ('mark', '>>>>>>> superhero'),
], title='No blank lines: one big conflict', file_title='monster-noblank/',
    notes=[
        'B1, the answer to the T14 vote. Show it right after the vote; don\'t run it live.',
        'Same two branches as monster/, but without the blank lines between panels. Changed '
        'lines that touch form one hunk, so all three lines conflict.',
        '"False alarm: nobody disagreed about FACE or LEGS. Git merges lines, not panels."',
    ])

# T15 · TERMINAL split: clean merge, broken program
sl, gl, gr = terminal_split(prs, 'Clean merge, broken program', [
    ('cmd', 'git switch rename'),
    ('cmd', 'git merge caller'),
    'Auto-merging m.py',
    "Merge made by the 'ort' strategy.",
    ('cmd', 'python3 m.py'),
    ('err', "NameError: name 'draw' is not defined"),
], [
    [('out', 'def '), ('hi', 'render'), ('out', '():')],
    '    print("monster")',
    '',
    'def main():',
    [('out', '    '), ('hi', 'draw()')],
    '',
    'main()',
], right_title='m.py after the merge', widths=(6.85, 4.60),
    pause='Git said no conflict. Who should have caught this?',
    notes=[
        'Live: cd ../broken && git merge caller && python3 m.py',
        'Branch `rename` changed `def draw():` to `def render():`. Branch `caller` changed '
        'main() to call draw(). Different lines, so no conflict. Verified on git 2.34.1.',
        'Pause answer: tests or CI after the merge; semantic or structure-aware merge tools.',
        'SAY: "Merging by lines is a design choice. It gives you false alarms AND silent '
        'misses." (T14 was the false alarm.)',
    ])

# T16 · BREAK
break_slide(prs, 3, notes=[
    'Labs put their Part-1 cards in their envelope ("that was a different repo").',
    'Ananya tapes the official #1-#4 on the Wall, puts `main` on #4, sets up the gc bin '
    '("unreachable · emptied by gc"), and puts a "You were cloned" copy pack on each table '
    '(#1-#4 with yellow `main` + blue `origin/main` on #4).',
    'Courier = the person nearest the Wall; same Courier all session.',
    'The break must start by 0:33. If behind, use the cut order in the cheat sheet.',
])

# =============================================================== PART 2 · Many tables
divider(prs, 'Many tables', BLUE, kicker='PART 2',
        notes=[
            '(5 seconds) "Part 2: many tables, one wall."',
            'Point at the Wall poster: "THE WALL · GitHub".',
        ])

# T18 · TALK git clone (Step 4, talk only)
sl = talk(prs, 'git clone', "Every table gets every card. Blue flag = where the wall's main was.",
          kicker='You were cloned', pause='We gave you all four cards, not just #4. Why?',
          step=4, msize=BAND_MEANING,
          notes=[
              'SAY: "While you were out, you were cloned. All three labs now build THE monster '
              'on the wall, from your own tables. Only your Courier may visit the wall."',
              '30-s check. Answer: the next merge needs the shared ancestor; log and history '
              'work offline.',
              'Then: "And the blue flag?" (your table\'s memory of where the wall\'s `main` '
              'was; it only moves when you go and look)',
              '"Every clone is a full repository. `origin` is a nickname for the Wall. GitHub '
              'is special only because we agreed it is."',
              '"From now on only your Courier visits the wall."',
          ])
sc = band_scene(sl, *band_bounds(sl))
sw, sh = stack_size()
for key in ('left', 'wall', 'right'):
    zz = sc[key]
    kind = 'branch' if key == 'wall' else 'origin'
    name = 'main' if key == 'wall' else 'origin/main'
    fsz = 14 if key == 'wall' else 13
    fw = chip_w(name, fsz, True, pad=0.12)
    grp = sw + fw - 0.15
    x0 = zz['l'] + (zz['w'] - grp) / 2
    st = card_stack(sl, x0, zz['t'] + 0.38, FOUR, s=S_BAND, dy=STACK_DY)
    flag(sl, st, name, kind=kind, edge='right', size=fsz)
ya = sc['wall']['t'] + 0.95
for (xa, xb), toward in ((sc['gap_l'], 'left'), (sc['gap_r'], 'right')):
    if toward == 'left':
        fetch_arrow(sl, xb - 0.12, ya, xa + 0.12, ya)
    else:
        fetch_arrow(sl, xa + 0.12, ya, xb - 0.12, ya)
    text(sl, xa, ya - 0.42, xb - xa, 0.32, 'git clone', size=16, bold=True, color=BLUE,
         font=MONO, align='c', anchor='m', markup=False)

# T19 · Step 5
task_slide(prs, 5, 'Each lab: add ONE card on top of #4.',
     ['Follow your lab mission', 'Number it. Move your yellow `main`.',
      'Read your number out loud'], minutes='2:00',
     notes=[
         'Missions: Lab 1 FACE -> dragon; Lab 2 BODY -> disco suit; Lab 3 LEGS -> roller '
         'skates (each copies the other panels exactly).',
         'Everyone shouts "#5". Ask: "How can three tables pick names that never clash, '
         'without talking?"',
         'Praise lab+number, then push: "Could the name come from the card itself?" Hand out '
         'the ID-rule slips.',
         'Expected IDs: Lab 1 drt3, Lab 2 cdt5, Lab 3 crr3.',
         'Wrong answer, random IDs: "Could two tables agree on the same card\'s ID without '
         'talking? Could anyone tell if a card was altered?"',
         'Check: "Lab 1 secretly changes LEGS to wheels. New ID? What happens to every card '
         'built on it?" (drw0; every later digit changes)',
         'Bonus: "dragon -> donkey?" (drt3 again: a collision)',
     ])

# T20 · TERMINAL split: hash IDs
sl = blank(prs)
text(sl, M, 0.55, CW, 0.80, 'The ID comes from the card itself', size=36, bold=True, anchor='m')
# Real output (git 2.34.1) of Lab 1's dragon commit made on top of the T13 merge in monster/
# (user.name "Lab 1"): the same object as the card on the right. The tree ID depends only on
# the content, so 5a5fee2 is the same on every machine; the parent ID and times differ.
term_panel(sl, M, 1.60, 6.80, 4.40, [
    ('cmd', 'git cat-file -p HEAD'),
    'tree 5a5fee2…',
    'parent 9f5707b…',
    'author Lab 1 <…> 1791104200 -0700',
    'committer Lab 1 <…> 1791104200 -0700',
    '',
    'face: dragon',
], size=20)
zx0, zx1 = 7.75, R
g = card(sl, (zx0 + zx1) / 2 - CARD_W / 2, 1.70, 'drt3', ('dragon', 'cape', 'tentacles'),
         parent='#4', author='Lab 1', s=1.0)
yb = underbrace(sl, g)
text(sl, zx0, yb + 0.10, zx1 - zx0, 0.42, 'ID = hash(everything on the card)', size=20,
     bold=True, align='c', anchor='m')
text(sl, zx0, yb + 0.60, zx1 - zx0, 0.80,
     'Ours: initials + check digit.\nGit: SHA-1 of the full commit text.', size=20,
     color=INK2, align='c', ls=1.10)
pause_bar(sl, 'Two labs draw identical cards on the same parent. Same ID?')
_set_notes(sl, [
    '"A commit is a tiny text object: a tree, its parents\' IDs, author, committer, message. '
    'Its ID is the hash of that text. No central counter, and history is tamper-evident: a '
    'changed card is literally a different card."',
    'Left and right are the SAME object: Lab 1\'s dragon card on top of #4 (the T13 merge). '
    'tree = the three panels, parent = "came from #4", author = Lab 1. Map them out loud.',
    'Optional LIVE (15 s, in monster/ right after the T13 merge): sed -i \'s/face: cat/face: '
    'dragon/\' monster.txt && git commit -qam "face: dragon" && git cat-file -p HEAD. The tree '
    'line (5a5fee2...) matches the slide on any machine; the parent ID, author and times '
    'differ.',
    'One sentence first: "At one table you could just agree on #2 and #3. Across three '
    'tables you can\'t, so the ID has to come from the card."',
    'LIVE, 20 s: echo \'face: smiley\' | git hash-object --stdin  on Shubham\'s laptop AND on '
    'one student\'s laptop. Both print 946ac5d... (verified). Same content, same ID, anywhere, '
    'no coordination.',
    'Pause answer: in our toy, yes (our rule ignores the author). In Git: same TREE, '
    'different COMMIT, because author and time are hashed too. (Model Critics: write that '
    'down.)',
    'Collision aside: Git uses 160 bits; SHAttered (2017) is why Git is moving toward '
    'SHA-256.',
])

# T21 · TERMINAL: tree of blobs
sl = blank(prs)
text(sl, M, 0.55, CW, 0.80, 'Inside a card: a tree of blobs', size=36, bold=True, anchor='m')
term_panel(sl, M, 1.60, CW, 3.90, [
    ('cmd', 'git ls-tree main'),
    [('out', '100644 blob d591a27…   body.txt')],
    [('out', '100644 blob 875229f…   face.txt')],
    [('out', '100644 blob '), ('hi', 'a27e19c…'), ('out', '   legs.txt')],
    ('cmd', 'git ls-tree cat-robot'),
    [('out', '100644 blob ba63925…   body.txt')],
    [('out', '100644 blob ef07ddc…   face.txt')],
    [('out', '100644 blob '), ('hi', 'a27e19c…'), ('out', '   legs.txt')],
], size=20)
text(sl, M, 5.60, CW, 0.40, 'card = commit  ·  panel = file (blob)  ·  all three panels = tree',
     size=20, color=INK2, anchor='m')
pause_bar(sl, 'LEGS never changed. How many copies of it does Git store?')
_set_notes(sl, [
    'Live in monster-files/ (one file per panel): git ls-tree main && git ls-tree cat-robot',
    'Pause answer: ONE. Same content -> same blob ID -> stored once. That\'s why "never '
    'erase" is cheap, and why writing "same" in an unchanged panel was blob reuse.',
    'Optional: git cat-file -p cat-robot^{tree} shows the same three lines: a tree is just a '
    'list of names and blob IDs.',
    'These blob IDs depend only on the content, so they match on every laptop. Commit IDs '
    'don\'t.',
    'Merging this 3-file repo also conflicts only on body.txt (verified).',
    'If behind: cut this slide to one sentence said over T20.',
])

# T22 · Step 6
task_slide(prs, 6, 'Courier: put your card on the wall. Move `main` to it.',
     ['First come, first served', "Then walk back from `main`. Is every lab's card there?"],
     minutes='6:00',
     notes=[
         'Ananya (the Wall) accepts the first two cards and moves `main` each time. Freeze: '
         '"Walk back from main. Lab 1, where\'s your dragon?" It isn\'t in main\'s history.',
         'The class invents the rule. Ananya moves `main` back to the first card, hands back '
         'the second ("REJECTED, fetch first"), and reads GUARD aloud.',
         'Check: "Your blue flag says #4. Where is the wall\'s `main` right now? How would you '
         'know?" (you can\'t, until you go and look)',
         'Rejected Couriers copy the new cards and move the blue flag (= fetch). Labs draw a '
         'combine card with two came-froms, compute its ID, move yellow `main` (= merge), and '
         're-post. Every combine merges automatically (different panels).',
         'AUDIT run 1 (30 s): Lab 1 traces how the tentacles reached `main`. They succeed '
         '(#3 Grace -> #4 Junio\'s merge -> lab + combine cards -> main).',
         'Wrong answers: one-parent redraw on top -> "you invented rebase"; Courier moves '
         '`main` anyway -> "that\'s push --force. Whose card just fell off?"',
     ])

# T23 · TALK git push / pull
sl = talk(prs, 'git push', 'The wall only moves forward. Your card must lead back to its main.',
          pause='Why refuse your card instead of just adding it?', step=6,
          msize=BAND_MEANING,
          notes=[
              'In the picture: Lab 3\'s crr3 reached the wall first. Lab 1\'s push of drt3 is '
              'refused. Lab 2 fetches: it copies crr3 and moves its blue flag. (Lab 2 still '
              'has its own cdt5 too; next it draws a combine card = the merge half of pull.)',
              '"Remotes accept only fast-forwards. Otherwise someone\'s card falls off `main`, '
              'which is what you saw."',
              '`git status` before fetching says "Your branch is ahead of \'origin/main\' by 1 '
              'commit." It only compares with your memory. After `git fetch`: "Your branch and '
              '\'origin/main\' have diverged, and have 1 and 1 different commits each" (both '
              'verified).',
              '"`pull --rebase` redraws your cards on top instead. Keep that in mind for Part '
              '3."',
              'On git 2.34 (tested) a plain `pull` of diverged branches refuses until you '
              'choose, so demo with --no-rebase (hidden backup B2). Thursday comes back to '
              'this.',
          ])
sc = band_scene(sl, *band_bounds(sl))
heads = {}
for key, nxt in (('left', DRT3), ('wall', CRR3), ('right', CRR3)):
    zz = sc[key]
    fname, fsz = {'left': ('', 13), 'wall': ('main', 14), 'right': ('origin/main', 13)}[key]
    over = max(0.0, (chip_w(fname, fsz, True, pad=0.12) - CARD_W * S_BAND) / 2) if fname else 0
    grp = sw + 0.32 + CARD_W * S_BAND + over
    x0 = zz['l'] + (zz['w'] - grp) / 2
    st = card_stack(sl, x0, zz['t'] + 0.38, FOUR, s=S_BAND, dy=STACK_DY)
    gn = card(sl, st['r'] + 0.32, st['t'], s=S_BAND, **nxt)
    link(sl, gn, st)
    heads[key] = gn
flag(sl, heads['wall'], 'main', size=14)
origin_flag(sl, heads['right'], size=13)
ya = sc['wall']['t'] + 0.72
xa, xb = sc['gap_l']
push_arrow(sl, xa + 0.12, ya, xb - 0.12, ya)
badge(sl, (xa + xb) / 2, ya, d=0.40, size=17)
text(sl, xa, ya + 0.32, xb - xa, 0.95, '! [rejected]\nmain -> main\n(fetch first)', size=16,
     color=RED, font=MONO, align='c', pitch=20, markup=False)
xa, xb = sc['gap_r']
fetch_arrow(sl, xa + 0.12, ya, xb - 0.12, ya)
text(sl, xa, ya - 0.42, xb - xa, 0.32, 'git fetch', size=16, bold=True, color=BLUE, font=MONO,
     align='c', anchor='m', markup=False)
text(sl, xa, ya + 0.22, xb - xa, 0.62, 'pull =\nfetch + merge', size=16, color=INK2,
     align='c', ls=1.1)

# =============================================================== PART 3 · Rewriting the past
divider(prs, 'Rewriting the past', ORANGE, kicker='PART 3',
        notes=[
            '(5 seconds) "Part 3: rewriting the past."',
            'BEFORE the next slide: Ananya fills in the mustache card (by Shubham, came from the '
            'current `main`, BODY disco suit, LEGS roller skates, ID computed: mdr1 if main is '
            'ddr2, mdr9 if ddr0), posts it, moves `main` to it, and says "New card on the wall. '
            'Everyone pull." Each Courier copies it.',
        ])

# T25 · Step 7
task_slide(prs, 7, 'The Client hates the mustache. Make it gone. For good.',
     ['Every lab already copied it', "Don't break anyone's table"], minutes='3:00',
     notes=[
         'Before this slide: Ananya has posted the mustache card and said "Everyone pull."',
         'Most labs rip it off (Ananya moves `main` back). Then: "Lab 3: add one card on your '
         'newest card and push." GUARD accepts it (walking back passes the mustache and '
         'reaches main), and the mustache comes back as a ZOMBIE. Real git does the same on a '
         'plain fast-forward push (verified).',
         'The invention we want: a FIX card that undoes it (dragon without the mustache), '
         'coming from the mustache card. Same picture as the card before, different ID.',
         'Nudge: "Lab 3 still holds that card. Then what?"',
         'Check: "When is ripping it off OK?" (only if nobody else has copied it yet)',
     ])

# T26 · TALK git revert vs git reset (custom two halves, 2-line meaning at 32 pt)
sl = blank(prs)
text(sl, M, TOP, CW, 0.50, 'You just invented', size=24, color=INK2, anchor='m')
text(sl, M, 1.12, CW, 1.05, 'git revert', size=54, bold=True, color=PURPLE, font=MONO,
     anchor='m', markup=False)
text(sl, M, 2.30, CW, 1.32, ['Shared? Add a fix card.',
                             'Private? You may move the flag back (`reset`).'],
     size=32, ls=1.12)
ptop = pause_bar(sl, 'Where did the ripped-off card go? Can you get it back?')
progress(sl, 7, y=7.12)
ly, bottom = 3.66, ptop - 0.18
hairline(sl, 6.67, ly, 6.67, bottom)
text(sl, M, ly, 5.65, 0.40, 'revert · safe when shared', size=20, bold=True, color=GREEN,
     anchor='m')
text(sl, 6.95, ly, R - 6.95, 0.40, 'reset · only if nobody copied it', size=20, bold=True,
     color=ORANGE, anchor='m')
s = 0.62                                    # bigger cards: mdr1 vs ddr8 must read
cyt = bottom - CARD_H * s - 0.02
gap_c = 0.50
cwid = CARD_W * s
lx0 = M + (5.65 - (3 * cwid + 2 * gap_c)) / 2
# ddr2 <- cdr2 matches T23's picture (crr3 reached the wall first, so the first combine is
# Lab 2's cdt5 + crr3 = cdr2) and Thursday's H4.
DDR2 = dict(cid='ddr2', panels=('dragon', 'disco', 'skates'), parent='cdr2')
MDR1 = dict(cid='mdr1', panels=('mustache', 'disco', 'skates'), parent='ddr2')
DDR8 = dict(cid='ddr8', panels=('dragon', 'disco', 'skates'), parent='mdr1')
a = card(sl, lx0, cyt, s=s, **DDR2)
b = card(sl, lx0 + cwid + gap_c, cyt, s=s, marks={0: 'changed'}, **MDR1)
c = card(sl, lx0 + 2 * (cwid + gap_c), cyt, s=s, marks={0: 'changed'}, **DDR8)
link(sl, b, a)
link(sl, c, b)
flag(sl, c, 'main', size=14)
lab_w = 2.05
rx0 = 6.95 + (R - 6.95 - (2 * cwid + gap_c + 0.22 + lab_w)) / 2
a2 = card(sl, rx0, cyt, s=s, **DDR2)
b2 = card(sl, rx0 + cwid + gap_c, cyt, s=s, state='ghost', cid='mdr1', parent='ddr2')
link(sl, b2, a2, color=GHOST, dash=True)
flag(sl, a2, 'main', size=14)
text(sl, b2['r'] + 0.22, b2['cy'] - 0.42, lab_w, 0.84, 'unreachable,\nnot deleted', size=20,
     bold=True, color=ORANGE, ls=1.1, anchor='m')
_set_notes(sl, [
    '"`git revert` adds an inverse commit (the fix card): new card, new ID, history intact. '
    'Safe when the commit is shared."',
    '"`git reset --hard` moves the branch flag back; the commit still exists, it is just '
    'unreachable. On a shared branch you\'d also need `push --force`, and you\'ve just seen '
    'what that does to everyone else (the zombie)."',
    'Pause answer: in the gc bin. "`git reflog` is your local diary of where HEAD has been. '
    'It finds that card, and the detached-HEAD card from T10, until gc prunes it. Defaults: '
    'unreachable reflog entries expire after 30 days, others after 90."',
    'IDs here follow the cheat sheet (main was ddr2 -> mustache mdr1 -> fix ddr8). If your '
    'wall ended on ddr0, say "mdr9 / ddr6" instead.',
    'The slide shows the textbook revert, drawn straight on the mustache card. LIVE, after the '
    'zombie, the wall\'s `main` is Lab 3\'s card (on top of the mustache), so the fix card must '
    'come from THAT card or GUARD rejects it. It undoes only the mustache and keeps Lab 3\'s '
    'change; its ID is whatever the rule gives (Model Critics check it).',
])

# T27 · Step 8
task_slide(prs, 8, 'Boss: "This history is a mess.\nOne clean card. Bin the rest."',
     ['Lab 3: open BOSS', 'Lab 1: open AUDIT again', 'Lab 2: does your newest card still fit?'],
     minutes='3:00',
     notes=[
         'Shubham plays the Boss STRAIGHT (the Boss has a point). Read the reason line aloud: '
         '"Reviewers can\'t read 14 combine cards; bisect needs every card on main to work."',
         'Lab 3 draws one clean card (came from: —, ID ddr7). Walking back never reaches '
         '`main`, so the Wall rejects it. Boss: "Tell the wall: FORCE."',
         'Ananya bins the old cards, then empties the bin, theatrically.',
         'AUDIT run 2: Lab 1 tries all three questions with the same card and the same wall. '
         'They fail. Ask: "Who lost what?"',
         'Check: "Does Lab 2\'s newest card still fit on the wall?" (No: it points at cards '
         'that are gone. They must redo their work on top of the clean card.)',
     ])

# T28 · TALK rebase / squash / force-push
sl = talk(prs, 'git push --force',
          'Rebase and squash draw NEW cards. --force makes the wall forget the old ones.',
          pause='The wall forgot. Where in this room does the truth still exist?', step=8,
          notes=[
              '"`git rebase -i` (squash) can\'t edit commits, because IDs are hashes. It '
              'writes new ones and moves the flag. `--force` moves the remote\'s flag."',
              'The six grey cards are the old wall history the audit needed: #3 (Grace drew '
              'the tentacles first), #4 (Junio settled BODY), the lab and combine cards, the '
              'mustache and its fix.',
              'Pause answer: Lab 2\'s stale table still holds every old card; local reflogs '
              'too. "Lab 1, you may now ask Lab 2\'s table. 30 seconds." They can answer again.',
              'Then: "Which copy does an auditor or a researcher mine? How long do table '
              'copies last?" (The server. Table copies last until gc or reflog expiry, and only '
              'on that laptop.)',
              '"The route the tentacles took to `main` is gone from the server. That\'s '
              'Tuesday\'s paper."',
          ])
z = sl.zone
s = 0.56
gw = CARD_W * s
gp = 0.40
row_w = 6 * gw + 5 * gp
ghost_ids = [('#3', None), ('#4', None), ('drt3', None), ('cdr2', None), ('mdr1', None),
             ('ddr8', None)]
wall_t = z['t'] + 0.14
gy = z['b'] - 0.16 - CARD_H * s
text(sl, M, gy - 0.50, row_w, 0.40, 'rewritten away', size=20, bold=True, color=ORANGE,
     align='c', anchor='m')
prev = None
for i, (cid, _) in enumerate(ghost_ids):
    gg = card(sl, M + i * (gw + gp), gy, cid, state='ghost', s=s)
    if prev:
        link(sl, gg, prev, color=GHOST, dash=True)
    prev = gg
wl = R - 2.80                    # narrowest wall that still fits its title pill (2.43 in)
wz = wall(sl, wl, wall_t, R - wl, z['b'] - wall_t)
# squash: the six old cards became one new card on the wall (orange = rewriting)
ay = gy + CARD_H * s / 2
xa0, xa1 = prev['r'] + 0.16, wl - 0.16
arrow(sl, xa0, ay, xa1, ay, color=ORANGE, w=3)
text(sl, xa0 - 0.2, ay - 0.44, xa1 - xa0 + 0.4, 0.34, 'squash', size=16, bold=True,
     color=ORANGE, font=MONO, align='c', anchor='m', markup=False)
# ddr7 + its flag stack ('main' with '--force' attached underneath), centred in the wall
DDR7 = dict(cid='ddr7', panels=('dragon', 'disco', 'skates'), parent='—')
fwf = chip_w('--force', 14, True, MONO, pad=0.12)
fwm = max(1.20 * s, chip_w('main', 14, True, pad=0.12))
grp = gw - 0.15 + max(fwm, fwf)
gnew = card(sl, wz['l'] + (wz['w'] - grp) / 2, gy, s=s, **DDR7)
fm = flag(sl, gnew, 'main', edge='right', size=14)
chip(sl, fm['l'], fm['b'] + 0.06, fwf, fm['h'], '--force', fill=ORANGE, color=INK, size=14,
     font=MONO, radius=0.04, rot=-4)

# T29 · PAUSE think-pair-share with contrasting cases (adds the "Does your rule allow…" lead)
sl = blank(prs)
Y29 = 0.25                                   # nudge the whole block toward the centre
text(sl, M, 0.40 + Y29, CW, 0.80, '💬', size=48, align='c', anchor='m', markup=False)
q29 = 'The Boss has a real reason. So does the auditor. Invent ONE wall rule that serves both.'
qsz = fit_size(q29, CW, 1.45, (40, 36, 32), bold=True, ls=1.06)
text(sl, M, 1.22 + Y29, CW, 1.45, q29, size=qsz, bold=True, align='c', anchor='m', ls=1.06,
     balance=True)
cases = ['squashing your 5 messy cards before anyone copied them?',
         'force-pushing `main` to drop the mustache after everyone pulled?',
         '"Squash and merge" of a teammate\'s 6-card idea via the PR button?']
mk_w = 0.62
bw = mk_w + max(line_width(cc, 24) for cc in cases) + 0.15
bx = (W - bw) / 2
text(sl, bx, 2.92 + Y29, bw, 0.42, 'Does your rule allow…', size=24, color=INK2, anchor='m')
rows(sl, bx, 3.46 + Y29, bw, cases, marks=['a ·', 'b ·', 'c ·'], size=24, mark_w=mk_w, gap=0.16)
rh = 'Think 1 min  ·  Pair 2 min  ·  Share 2 min'
cw_ = max(6.20, chip_w(rh, 20))
chip(sl, (W - cw_) / 2, 5.50 + Y29, cw_, 0.56, rh, fill=LAVENDER, color=INK, size=20)
_set_notes(sl, [
    'Think 1 min (alone) · Pair 2 min · Share 2 min. This is "inventing with contrasting '
    'cases": students invent the rule first; the next slide is the official answer.',
    'Expected sort: (a) OK, it\'s private; (b) not OK: it\'s shared, it causes zombies, and it '
    'hides history; (c) the interesting one.',
    'Plan for wrong answers: a rule that only bans force-push lets (c) through: "No force '
    'needed. What did the auditor lose?" (the 6 cards and their route to main). (c) matches '
    'the squash example in Just et al.\'s abstract; no numbers.',
    'A pair proposes "never rewrite main": hand them the KEY card ("The mustache card had the '
    'Client\'s API key written on it. Is a fix card enough?"). No: the key is still in every '
    'clone. Rewrite AND rotate the key.',
    'Write their rules on the board. Rules to expect or seed: protected main; squash only '
    'cards nobody else has; keep merge cards; tag or archive the branch before rewriting; '
    '--force-with-lease; the wall logs every push.',
])

# T30 · PAPER
sl = paper(prs, 'Switching to Git: the Good, the Bad, and the Ugly',
      'Just, Herzig, Czerwonka, Murphy · ISSRE 2016 · Microsoft Research',
      [('GOOD', GREEN, 'Cheap branches.\nLocal commits.', None),
       ('BAD', INK, 'Many routes to main.\nHard to trace.', None),
       ('UGLY', ORANGE, 'Rebase, squash,\nforce-push\nerase history.', None)],
      bottom='Their algorithm: rebuild **integration paths** from what\'s left.',
      notes=[
          '"The Good/Bad/Ugly split is our reading of the abstract." Don\'t quote numbers, and '
          'don\'t explain how their algorithm works: we only have the abstract.',
          '"An integration path is the route a change took to `main`. You traced one in Step 6 '
          'and couldn\'t in Step 8. Auditors (e.g., the history of security or privacy code) '
          'and researchers mining history depend on it."',
          '"Compare their countermeasures with your wall rules."',
          '"Research angle: version history is a dataset that users can silently delete."',
      ])
for tag, x in zip(('Steps 2–3', 'Step 6', 'Step 8'), (0.75, 4.78, 8.80)):   # paper() columns
    text(sl, x, 4.92, 3.78, 0.42, tag, size=20, color=INK2, anchor='m')


def exit_ticket(prs, title, sub, code='____', qr=None, notes=None):
    """MENTI slide without options: same chip / QR / code geometry as style.menti(), with a
    44 pt title + 26 pt INK2 subline top-aligned to the QR box (y 1.40)."""
    sl = blank(prs)
    menti_chip(sl)
    text(sl, M, 1.40 - 0.13, 8.10, 0.80, title, size=44, bold=True, anchor='t', ls=1.0)
    text(sl, M, 2.22, 8.10, 0.50, sub, size=26, color=INK2, anchor='t')
    if qr and os.path.exists(qr):
        sl.shapes.add_picture(qr, Inches(9.333), Inches(1.40), Inches(3.25), Inches(3.25))
    else:
        b = box(sl, 9.333, 1.40, 3.25, 3.25, fill=WHITE, line='BBBBBB', lw=2, dash=True,
                radius=0.12)
        label(b, 'QR', size=24, bold=True, color=MUTED)
    chip(sl, 9.333, 4.85, 3.25, 0.60, f'menti.com  ·  {code}', size=18)
    return _set_notes(sl, notes)


# T31 · MENTI exit ticket
exit_ticket(prs, 'Exit ticket', '4 questions  ·  laptops closed', code=MENTI_CODE, qr=MENTI_QR,
      notes=[
          'Questions are in spec section 8 / menti_questions.md (Q-a to Q-d). Laptops closed.',
          'Q-a tests the shared ancestor (answer: Nothing), Q-b that cards never change '
          '(changes, and the old commit still exists), Q-c fixing a rejected push (git pull, '
          'then git push), Q-d revert vs reset (git revert, then push).',
          'Read the results out. ANY item under 70% correct: redo it on cards right away (2 '
          'min from the buffer), and open Thursday\'s H4 with it.',
      ])

# T32 · RECAP
recap(prs, [("Can't go back", 'Never erase: new card', 'git commit'),
            ('Shuffled pile', '"came from __"', 'git log'),
            ('Risky ideas', 'flags; you are here', 'branch · HEAD'),
            ('Two ideas, one monster', 'card with two parents', 'git merge'),
            ('Work from my table', 'copy every card', 'clone · origin/main'),
            ('Two #5s', 'ID from the card itself', 'SHA-1 hash'),
            ('My card fell off', 'look, combine, post', 'fetch · pull · push'),
            ('Bad card, already copied', 'fix card', 'git revert'),
            ('Boss wants it neat', 'bin + FORCE', 'rebase · --force')],
      notes=[
          '"Almost every Git command makes a card, moves a flag, or copies cards between '
          'tables."',
          '"Now: where does our model lie?" (next slide)',
      ])

# T33 · PAUSE Model Critics
pause(prs, 'Name one place our paper model lies.', rhythm='Model Critics  ·  30 s each',
      notes=[
          'Each Model Critic reads one entry from "Where does paper ≠ Git?" (3 x 30 s).',
          'Seed answers, if needed:',
          '- Our card is a snapshot; `git gc` packs objects as deltas (packfiles). Both true: '
          'snapshots in the model, deltas on disk.',
          '- The index is a real file (.git/index); our cards have no equivalent.',
          '- Git merges lines, not panels (T14, T15).',
          '- Our hash is a toy: it ignores author and time, and it collides (donkey).',
          '- origin/main moves only on fetch or push.',
          '- The reflog is local and never pushed.',
          '- Two pairs shared one table; in Git each clone or worktree has exactly one HEAD.',
          'Bridge: "A model is a claim too. On Thursday we\'ll check other people\'s claims."',
      ])

# T34 · STARRY
starry(prs, 'Cards never change.\nFlags move.\nWalls copy cards.',
       small='…and Git never erases, until you force it to.',
       notes=[
           '"Thursday: from how Git works to how humans struggle with it."',
           'Homework: "Read the Stack Overflow study (Yang et al., TOSEM 2022) and bring one '
           'claim you don\'t believe, with its page number. You\'ll write it on a sticky at the '
           'door."',
           'Buffer (3 min): redo any exit item under 70% on cards now.',
       ])

# =============================================================== hidden backups (B1 is slide 15)
sl, _ = terminal(prs, [
    ('cmd', 'git push'),
    ('err', ' ! [rejected]        main -> main (fetch first)'),
    ('cmd', 'git fetch'),
    ('cmd', 'git status'),
    "Your branch and 'origin/main' have diverged,",
    ('cmd', 'git pull --no-rebase'),
    ('ok', "Merge made by the 'ort' strategy."),
    ('cmd', 'git push'),
    ('ok', '   …  main -> main'),
], title='Step 6 in real Git',
    notes=[
        'BACKUP B2 (hidden). Optional Step-6 rehearsal (verified): a bare wall.git cloned 3 '
        'times, each clone changes a different line.',
        'The 2nd and 3rd pushes are rejected with "(fetch first)". git status says "ahead of '
        'origin/main by 1 commit" until you fetch; then "have diverged".',
        'On git 2.34 a plain pull of diverged branches refuses until you choose, hence '
        '--no-rebase. Pushing again then succeeds.',
    ])
hide(sl)

sl = blank(prs)
text(sl, M, 0.45, CW, 0.85, 'Fast lab? Try one.', size=40, bold=True, anchor='m')
grid(sl, M, 1.55, [(3.40, 'cmd'), (8.23, 'ink')], ('Git', 'On paper'), [
    ('git tag v1.0', 'A flag that never moves.'),
    ('git stash', 'Park a half-drawn card; take it back later.'),
    ('git cherry-pick', 'Copy ONE change onto your branch. New card, new ID.'),
    ('git blame', 'For each panel, who last changed it? Use only came-from.'),
], pitch=0.85, size=24)
_set_notes(sl, [
    'BACKUP B3 (hidden). Power-ups for fast labs (only if ahead of time).',
    'tag: a flag that never moves (releases). stash: park a half-drawn card. cherry-pick: '
    'copies one change as a NEW card with a NEW ID (different parent). blame: walk came-from '
    'arrows back, per panel.',
])
hide(sl)

# =============================================================== spec labels in the notes
# The cheat sheet and spec say "T15", but B1 is a normal slide after T14, so from T15 on the
# deck slide number is T-number + 1. Put the spec label + clock at the top of every note.
LABELS = [
    'T1 · TITLE · 0:00', 'T2 · MENTI hook · 0:00–0:02', 'T3 · DIVIDER Part 1 · 0:02',
    'T4 · TASK Step 0 · 0:02 (block 0:02–0:07)', 'T5 · TALK git commit · ~0:05',
    'T6 · GOALS · 0:07', 'T7 · TASK Step 1 · 0:08 (block 0:08–0:13)', 'T8 · TALK git log · ~0:12',
    'T9 · TASK Step 2 · 0:13 (block 0:13–0:20)', 'T10 · TALK git branch · ~0:18',
    'T11 · TASK Step 3 · 0:20 (block 0:20–0:28)', 'T12 · CONFLICT git merge · ~0:26',
    'T13 · TERMINAL real merge · 0:28 (block 0:28–0:33)', 'T14 · MENTI predict · ~0:30',
    'B1 · answer to T14 · ~0:31', 'T15 · TERMINAL clean merge, broken program · ~0:32',
    'T16 · BREAK · 0:33–0:36', 'T17 · DIVIDER Part 2 · 0:36',
    'T18 · TALK git clone (Step 4) · 0:36', 'T19 · TASK Step 5 · 0:37 (block 0:37–0:42)',
    'T20 · TERMINAL hash IDs · ~0:40', 'T21 · TERMINAL tree of blobs · ~0:41',
    'T22 · TASK Step 6 · 0:42 (block 0:42–0:51)', 'T23 · TALK git push / pull · ~0:49',
    'T24 · DIVIDER Part 3 · 0:51', 'T25 · TASK Step 7 · 0:51 (block 0:51–0:57)',
    'T26 · TALK revert vs reset · ~0:55', 'T27 · TASK Step 8 · 0:57 (block 0:57–1:02)',
    'T28 · TALK rebase / squash / force-push · ~1:00', 'T29 · PAUSE think-pair-share · 1:02–1:07',
    'T30 · PAPER Good / Bad / Ugly · 1:07–1:09', 'T31 · MENTI exit ticket · 1:09–1:13',
    'T32 · RECAP · 1:13', 'T33 · PAUSE Model Critics · 1:14–1:16', 'T34 · STARRY · 1:16',
    'B2 · hidden backup', 'B3 · hidden backup',
]
assert len(LABELS) == len(prs.slides), (len(LABELS), len(prs.slides))
for i_, (s_, lab_) in enumerate(zip(prs.slides, LABELS), 1):
    tf_ = s_.notes_slide.notes_text_frame
    tf_.text = f'{lab_} · deck slide {i_}\n\n{tf_.text}'

# =============================================================== save + render
save(prs, OUT)
print('slides:', len(prs.slides), '(hidden backups:', len(BACKUPS), ')')
if '_' in MENTI_CODE or not os.path.exists(MENTI_QR):
    print('REMINDER: Menti QR / access code still placeholders on T2, T14 and T31 '
          f'(set MENTI_CODE, save the QR to {MENTI_QR}, rebuild).')

if '--no-render' not in sys.argv:
    from pptx import Presentation as _P
    rv = _P(OUT)
    for s_ in rv.slides:
        if s_._element.get('show') is not None:
            del s_._element.attrib['show']
    rv.save(REVIEW)
    if os.path.isdir(REN):
        for f in os.listdir(REN):
            if f.endswith('.png'):
                os.remove(os.path.join(REN, f))
    os.makedirs(REN, exist_ok=True)
    shutil.rmtree('/tmp/lo_tuesday', ignore_errors=True)
    pngs = render(REVIEW, REN, 't', profile='lo_tuesday')
    print(len(pngs), 'PNGs ->', REN)
