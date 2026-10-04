"""CS294 Git Week -- visual system helpers (python-pptx 1.0.x, Google-Slides-safe).

All geometry is in inches on a 13.333 x 7.5 in (16:9) canvas.
Import:  sys.path.insert(0, '/home/shagarw_google_com/shubham/git_course/design'); from vis import *
"""
import random
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.oxml.ns import qn
from lxml import etree

# ---------------------------------------------------------------- tokens
W, H = 13.333, 7.5          # canvas
M = 0.75                    # side margin
CW = W - 2 * M              # 11.833 content width
TOP = 0.6                   # top of chips / kicker row

C = dict(
    BG='FFFFFF', INK='111111', INK2='555555', MUTED='919191', HAIR='E3E3E3',
    PURPLE='9437FF', PINK='FF2F92', BLUE='00A2FF', ORANGE='FF9300', LAV='F6ECF8',
    RED='E02424', RED_T='FDECEC', GREEN='1E9E4A', GREEN_T='E6F6EC',
    YELLOW='FFE45E', YELLOW_T='FFF6C7', BLUE_T='E8F6FF', PANEL='FAFAFA',
    CARD_LINE='1F1F1F', NAVY='0A1631', TERM='16181D', GHOST='BDBDBD', GHOST_T='F3F3F3',
)
FONT = 'Inter'
MONO = 'Noto Sans Mono'

ALIGN = dict(l=PP_ALIGN.LEFT, c=PP_ALIGN.CENTER, r=PP_ALIGN.RIGHT)
ANCH = dict(t=MSO_ANCHOR.TOP, m=MSO_ANCHOR.MIDDLE, b=MSO_ANCHOR.BOTTOM)


def rgb(h):
    return RGBColor.from_string(h)


# ---------------------------------------------------------------- deck / slide
def new_deck():
    prs = Presentation()
    prs.slide_width = Emu(12192000)   # 13.333 in
    prs.slide_height = Emu(6858000)   # 7.5 in
    return prs


def blank(prs, bg='FFFFFF'):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    f = s.background.fill
    f.solid()
    f.fore_color.rgb = rgb(bg)
    return s


def notes(slide, txt):
    slide.notes_slide.notes_text_frame.text = txt


# ---------------------------------------------------------------- low-level
def _no_shadow(shp):
    """Kill theme-driven shadow + white default text: drop <p:style>, add empty effectLst.
    (python-pptx autoshapes/connectors reference theme effectStyle 2 = outer shadow.)"""
    el = shp._element
    st = el.find(qn('p:style'))
    if st is not None:
        el.remove(st)
    spPr = el.spPr
    if spPr.find(qn('a:effectLst')) is None:
        ef = etree.SubElement(spPr, qn('a:effectLst'))
        # schema order: ... ln, effectLst, scene3d, sp3d, extLst
        for tag in ('a:scene3d', 'a:sp3d', 'a:extLst'):
            nxt = spPr.find(qn(tag))
            if nxt is not None:
                nxt.addprevious(ef); break


def box(sl, x, y, w, h, fill=None, line=None, lw=1.5, shape=MSO_SHAPE.RECTANGLE,
        radius=None, dash=False, rot=0):
    """radius in inches (absolute) for rounded rects; 'pill' => fully round."""
    if radius is not None and shape == MSO_SHAPE.RECTANGLE:
        shape = MSO_SHAPE.ROUNDED_RECTANGLE
    shp = sl.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        r = 0.5 if radius == 'pill' else min(0.5, (radius or 0.1) / min(w, h))
        shp.adjustments[0] = r
    if fill:
        shp.fill.solid(); shp.fill.fore_color.rgb = rgb(fill)
    else:
        shp.fill.background()
    if line:
        shp.line.color.rgb = rgb(line); shp.line.width = Pt(lw)
        if dash:
            shp.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    else:
        shp.line.fill.background()
    if rot:
        shp.rotation = rot
    _no_shadow(shp)
    tf = shp.text_frame
    for side in ('left', 'right', 'top', 'bottom'):
        setattr(tf, f'margin_{side}', Inches(0.04))
    return shp


def _fill_tf(tf, content, size, bold, color, align, font, italic, ls):
    """content: str (\n = new paragraph) or list of paragraphs; a paragraph is a str
    or a list of runs; a run is a str or (text, dict(size,bold,color,font,italic))."""
    if isinstance(content, str):
        content = content.split('\n')
    tf.clear()
    for i, para in enumerate(content):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = ALIGN[align]
        if ls:
            p.line_spacing = ls
        runs = [para] if isinstance(para, (str, tuple)) else para
        for run in runs:
            t, o = (run, {}) if isinstance(run, str) else run
            r = p.add_run(); r.text = t
            f = r.font
            f.name = o.get('font', font); f.size = Pt(o.get('size', size))
            f.bold = o.get('bold', bold); f.italic = o.get('italic', italic)
            f.color.rgb = rgb(o.get('color', color))


def text(sl, x, y, w, h, content, size=28, bold=False, color='111111', align='l',
         anchor='t', font=FONT, italic=False, ls=None, margin=0.0):
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = ANCH[anchor]
    for side in ('left', 'right', 'top', 'bottom'):
        setattr(tf, f'margin_{side}', Inches(margin))
    _fill_tf(tf, content, size, bold, color, align, font, italic, ls)
    return tb


def label(shp, content, size=16, bold=True, color='111111', align='c', anchor='m',
          font=FONT, italic=False, ls=None):
    tf = shp.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = ANCH[anchor]
    _fill_tf(tf, content, size, bold, color, align, font, italic, ls)
    return shp


def chip(sl, x, y, w, h, content, fill='111111', color='FFFFFF', size=18, bold=True,
         line=None, lw=1.5, font=FONT, radius='pill', rot=0):
    return label(box(sl, x, y, w, h, fill=fill, line=line, lw=lw, radius=radius, rot=rot),
                 content, size=size, bold=bold, color=color, font=font)


def arrow(sl, x1, y1, x2, y2, color='555555', w=2.0, dash=False, head=True):
    """Straight connector from (x1,y1) to (x2,y2); arrowhead at (x2,y2)."""
    c = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1),
                                Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color); c.line.width = Pt(w)
    if dash:
        c.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    ln = c.line._get_or_add_ln()
    if head:
        te = etree.SubElement(ln, qn('a:tailEnd')); te.set('type', 'triangle')
        te.set('w', 'med'); te.set('len', 'med')
    _no_shadow(c)
    return c


# ---------------------------------------------------------------- primitives
CARD_W, CARD_H = 1.70, 2.45
ROWS = ('HEAD', 'BODY', 'LEGS')


def card(sl, x, y, cid='a3f9', panels=('🙂', '📦', '🦵'), parent=None, author=None,
         s=1.0, state='committed', marks=None):
    """Commit card. state: committed | staged | sketch | ghost.
    marks: {row_index: 'changed'|'auto'|'conflict'} tints a panel.
    parent: str shown as '← parent' on the bottom line. Returns geometry dict."""
    marks = marks or {}
    w, h = CARD_W * s, CARD_H * s
    pad = 0.12 * s
    ghost = state == 'ghost'
    body_line = {'committed': C['CARD_LINE'], 'staged': C['CARD_LINE'],
                 'sketch': '9A9A9A', 'ghost': C['GHOST']}[state]
    box(sl, x, y, w, h, fill=(C['GHOST_T'] if ghost else 'FFFFFF'), line=body_line,
        lw=1.5, radius=0.12 * s, dash=state in ('sketch', 'ghost'))
    # ID tag
    tw, th = 0.86 * s, 0.34 * s
    if state == 'committed':
        chip(sl, x + pad, y + pad, tw, th, cid, fill=C['INK'], color='FFFFFF',
             size=max(11, 14 * s), font=MONO, radius=0.06 * s)
    elif state == 'ghost':
        chip(sl, x + pad, y + pad, tw, th, cid, fill=None, line=C['GHOST'], lw=1,
             color=C['GHOST'], size=max(11, 14 * s), font=MONO, radius=0.06 * s)
    else:  # staged / sketch: empty dashed tag = no ID yet
        chip(sl, x + pad, y + pad, tw, th, '', fill=None, line='9A9A9A', lw=1,
             radius=0.06 * s).line.dash_style = MSO_LINE_DASH_STYLE.DASH
    if author:
        text(sl, x + pad + tw, y + pad, w - 2 * pad - tw, th, author, size=max(11, 13 * s),
             color=(C['GHOST'] if ghost else C['INK2']), align='r', anchor='m')
    # panels
    ph, gap, py0 = 0.48 * s, 0.06 * s, y + 0.56 * s
    for i, emo in enumerate(panels):
        mk = marks.get(i)
        fill, ln, lw, dash = C['PANEL'], 'DDDDDD', 0.75, False
        if mk == 'changed':
            fill, ln = C['YELLOW_T'], 'E8C547'
        elif mk == 'auto':
            fill, ln, lw = C['GREEN_T'], C['GREEN'], 1.5
        elif mk == 'conflict':
            fill, ln, lw, dash = C['RED_T'], C['RED'], 2.25, True
        if ghost:
            fill, ln = C['GHOST_T'], 'DDDDDD'
        p = box(sl, x + pad, py0 + i * (ph + gap), w - 2 * pad, ph, fill=fill, line=ln,
                lw=lw, dash=dash, radius=0.05 * s)
        if not ghost:
            label(p, emo, size=26 * s, bold=False)
    if parent:
        text(sl, x + pad, y + 2.12 * s, w - 2 * pad, 0.26 * s,
             [[('← ', {'font': FONT}), (parent, {})]],
             size=max(11, 13 * s), color=(C['GHOST'] if ghost else C['INK2']),
             font=MONO, align='c', anchor='m')
    return dict(x=x, y=y, w=w, h=h, l=x, r=x + w, t=y, b=y + h, cx=x + w / 2, cy=y + h / 2,
                row_y=[py0 + i * (ph + gap) + ph / 2 for i in range(3)], s=s)


def row_labels(sl, g, size=14):
    """HEAD/BODY/LEGS labels to the left of a card geometry g."""
    for i, nm in enumerate(ROWS):
        text(sl, g['l'] - 0.85, g['row_y'][i] - 0.15, 0.75, 0.3, nm, size=size,
             bold=True, color=C['MUTED'], align='r', anchor='m')


def link(sl, child, parent, color='555555', w=2.0, dash=False):
    """Parent pointer: arrow from child's left edge to parent's right edge."""
    return arrow(sl, child['l'] - 0.04, child['cy'], parent['r'] + 0.04, parent['cy'],
                 color=color, w=w, dash=dash)


def flag(sl, g, name, slot=0, color=None):
    """Sticky-note branch flag stuck on the top edge of card g. slot shifts right."""
    s = g['s']; fw, fh = 1.2 * s, 0.5 * s
    fx = g['cx'] - fw / 2 + slot * (fw + 0.1 * s)
    fy = g['t'] - fh * 0.75
    chip(sl, fx, fy, fw, fh, name, fill=color or C['YELLOW'], color=C['INK'],
         size=max(12, 16 * s), radius=0.04, rot=-4)
    return dict(x=fx, y=fy, w=fw, h=fh, cx=fx + fw / 2)


def head_pin(sl, cx, top):
    """Pink HEAD pin centered at cx, its pointer tip touching y=top."""
    pw, ph = 0.86, 0.34
    tri = box(sl, cx - 0.09, top - 0.12, 0.18, 0.12, fill=C['PINK'],
              shape=MSO_SHAPE.ISOSCELES_TRIANGLE, rot=180)
    chip(sl, cx - pw / 2, top - 0.12 - ph, pw, ph, 'HEAD', fill=C['PINK'], size=13)
    return top - 0.12 - ph


def conflict_tag(sl, g, row=1):
    """Red 'CONFLICT' pill to the right of panel `row` of card g."""
    chip(sl, g['r'] + 0.12, g['row_y'][row] - 0.17, 1.25, 0.34, 'CONFLICT',
         fill=C['RED'], size=13)


def wall(sl, x, y, w, h, title='THE WALL  ·  GitHub'):
    box(sl, x, y, w, h, fill=C['BLUE_T'], line=C['BLUE'], lw=2, radius=0.18)
    chip(sl, x + 0.2, y - 0.2, 3.0, 0.4, title, fill=C['BLUE'], size=15)


def table_zone(sl, x, y, w, h, title):
    box(sl, x, y, w, h, fill=C['PANEL'], line='CCCCCC', lw=1.25, radius=0.14)
    text(sl, x + 0.2, y + 0.1, w - 0.4, 0.35, title, size=15, bold=True, color=C['INK2'])


def badge(sl, cx, cy, glyph='×', fill=None, d=0.42):
    chip(sl, cx - d / 2, cy - d / 2, d, d, glyph, fill=fill or C['RED'], size=16,
         radius='pill')


def progress(sl, step, total=9, y=6.92, on_dark=False):
    """Row of step dots bottom-right; current = purple, done = ink, future = hairline."""
    d, gap = 0.13, 0.11
    x0 = W - M - (total * d + (total - 1) * gap)
    for i in range(total):
        col = C['PURPLE'] if i == step else (C['INK'] if i < step else 'D6D6D6')
        if on_dark and i < step:
            col = 'FFFFFF'
        box(sl, x0 + i * (d + gap), y, d, d, fill=col, shape=MSO_SHAPE.OVAL)


# ---------------------------------------------------------------- layouts
def title_slide(prs, title, subtitle=None, who='Shubham & Ananya',
                topic='Git Week · Implementation Day'):
    sl = blank(prs)
    text(sl, M, 2.3, CW, 1.4, title, size=60, align='c', anchor='b', ls=1.0)
    if subtitle:
        text(sl, M, 3.85, CW, 0.7, subtitle, size=26, color=C['INK2'], align='c')
    g = C['MUTED']
    text(sl, 0.3, 6.2, 6.7, 0.55,
         [[('CS294: ', {}), ('Modern Programming Tools', {'bold': True})]],
         size=26, color=g, anchor='m')
    text(sl, 7.05, 6.2, 2.6, 0.55, 'UC Berkeley', size=26, color=g, align='c', anchor='m')
    text(sl, 9.65, 6.2, 3.38, 0.55, who, size=26, color=g, align='r', anchor='m')
    bar = box(sl, 0, 6.81, W, 0.48, fill=C['PURPLE'])
    text(sl, 0.3, 6.81, W - 0.6, 0.48, topic, size=26, italic=True, color='FFFFFF',
         align='r', anchor='m')
    return sl


ON_DARK = {C['PURPLE'], C['NAVY'], C['TERM'], C['INK']}


def divider(prs, title, color='FF2F92', kicker=None, emoji=None):
    """Full-bleed colour. Title centred on the slide's optical centre (box 3.0..4.5)."""
    sl = blank(prs, color)
    fg = 'FFFFFF' if color in ON_DARK else C['INK']
    if emoji:
        text(sl, M, 1.25, CW, 1.05, emoji, size=54, align='c', anchor='b')
    if kicker:
        text(sl, M, 2.45, CW, 0.5, kicker, size=24, bold=True, color=fg, align='c', anchor='b')
    text(sl, M, 3.0, CW, 1.5, title, size=60, color=fg, align='c', anchor='m', ls=1.0)
    return sl


def task(prs, step, instruction, lines=(), minutes=None, visual=False, step_label=None):
    """Lavender activity slide. lines: up to 3 short sub-steps. visual=True halves the
    text column (x 0.75..7.35) and leaves a visual zone x 7.75..12.58, y 1.55..6.45."""
    sl = blank(prs, C['LAV'])
    chip(sl, M, TOP, 1.55, 0.5, step_label or f'STEP {step}', size=18)
    if minutes:
        chip(sl, W - M - 1.75, TOP, 1.75, 0.5, f'⏱  {minutes} min', fill='FFFFFF',
             color=C['INK'], line=C['INK'], lw=1.5, size=20)
    tw = 6.6 if visual else CW
    text(sl, M, 1.55, tw, 2.25, instruction, size=44, bold=True, ls=1.05)
    if lines:
        paras = [[(f'{i+1}  ', {'bold': True, 'color': C['PURPLE']}), (ln, {})]
                 for i, ln in enumerate(lines)]
        text(sl, M, 4.0, tw, 2.45, paras, size=28, ls=1.25)
    if isinstance(step, int):
        progress(sl, step)
    return sl


def talk(prs, command, meaning, kicker='You just invented', pause=None, step=None):
    """Reveal slide after a task. meaning <= 2 lines (~70 chars).
    Diagram zone: x 0.75..12.58, y 3.75..6.0 with pause bar (..6.85 without); cards s<=0.9."""
    sl = blank(prs)
    text(sl, M, TOP, CW, 0.5, kicker, size=24, color=C['INK2'], anchor='m')
    text(sl, M, 1.12, CW, 1.05, command, size=54, bold=True, color=C['PURPLE'],
         font=MONO, anchor='m')
    text(sl, M, 2.3, CW, 1.2, meaning, size=32, ls=1.15)
    if pause:
        b = box(sl, M, 6.2, CW, 0.72, fill=C['LAV'], radius=0.14)
        text(sl, M + 0.3, 6.2, CW - 0.6, 0.72, '💬  ' + pause, size=24, anchor='m')
    if isinstance(step, int):
        progress(sl, step, y=7.12)
    return sl


def menti(prs, question, code='____ ____', options=None):
    sl = blank(prs)
    chip(sl, M, TOP, 2.2, 0.48, '📊  Live poll', fill=C['PURPLE'], size=18)
    text(sl, M, 1.4, 8.1, 3.3, question, size=44, bold=True, anchor='m', ls=1.08)
    qr = box(sl, 9.33, 1.4, 3.25, 3.25, fill='FFFFFF', line='BBBBBB', lw=2, dash=True,
             radius=0.12)
    label(qr, 'QR', size=24, bold=True, color=C['MUTED'])
    chip(sl, 9.33, 4.85, 3.25, 0.6, f'menti.com  ·  {code}', size=18)
    if options:
        paras = [[(f'{"ABCD"[i]}  ', {'bold': True, 'color': C['PURPLE']}), (o, {})]
                 for i, o in enumerate(options[:4])]
        text(sl, M, 4.95, 8.1, 1.9, paras, size=26, ls=1.15)
    return sl


TERM_COL = dict(cmd='F5F5F5', out='A9B1BD', err='FF6B6B', ok='5AD17A', mark='FFB020',
                dim='7C8594', hi='C792EA')


def terminal(prs_or_slide, lines, title=None, x=M, y=1.6, w=CW, h=5.1, size=20):
    """lines: list of (kind, text); kind in TERM_COL. 'cmd' gets a dim '$ ' prompt.
    Capacity (panel h=5.1): <=11 lines at 20pt, <=13 at 18pt (min size). ~66 chars/line at
    20pt full width; ~36 chars at 20pt in a half-width (w=5.8) panel."""
    sl = blank(prs_or_slide) if hasattr(prs_or_slide, 'slides') else prs_or_slide
    if title:
        text(sl, M, 0.55, CW, 0.8, title, size=36, bold=True, anchor='m')
    box(sl, x, y, w, h, fill=C['TERM'], radius=0.16)
    for i, col in enumerate(('FF5F57', 'FEBC2E', '28C840')):
        box(sl, x + 0.25 + i * 0.25, y + 0.2, 0.16, 0.16, fill=col, shape=MSO_SHAPE.OVAL)
    paras = []
    for kind, t in lines:
        if kind == 'cmd':
            paras.append([('$ ', {'color': TERM_COL['dim']}), (t, {'color': TERM_COL['cmd']})])
        else:
            paras.append([(t if t else ' ', {'color': TERM_COL[kind]})])
    text(sl, x + 0.3, y + 0.55, w - 0.6, h - 0.75, paras, size=size, font=MONO, ls=1.1,
         color=TERM_COL['out'])
    return sl


def statement(prs, content, sub=None, size=54):
    """content may be rich runs; colour ONE key phrase purple."""
    sl = blank(prs)
    text(sl, M, 1.6, CW, 3.9, content, size=size, bold=True, align='c', anchor='m', ls=1.1)
    if sub:
        text(sl, M, 5.7, CW, 0.6, sub, size=22, color=C['MUTED'], align='c')
    return sl


def starry(prs, statement_text, kicker="If you're going to remember one thing…", seed=294):
    sl = blank(prs, C['NAVY'])
    rnd = random.Random(seed)
    # horizon glow + hills (drawn first so stars sit above the glow)
    box(sl, 2.6, 5.35, 8.2, 1.3, fill='12345A', shape=MSO_SHAPE.OVAL)
    box(sl, -2.5, 5.75, 10.5, 4.2, fill='071021', shape=MSO_SHAPE.OVAL)
    box(sl, 5.0, 5.55, 11.0, 4.6, fill='050B18', shape=MSO_SHAPE.OVAL)
    for _ in range(95):
        d = rnd.choice((0.025, 0.03, 0.035, 0.045, 0.06))
        col = rnd.choices(('FFFFFF', 'CADCFF', '7F94B8'), weights=(4, 3.5, 2.5))[0]
        sx, sy = rnd.uniform(0.1, W - 0.1), rnd.uniform(0.08, 5.5)
        box(sl, sx, sy, d, d, fill=col, shape=MSO_SHAPE.OVAL)
    text(sl, M, 0.6, CW, 0.9, kicker, size=36, color='FFFFFF', align='c', anchor='m')
    text(sl, M, 1.85, CW, 3.0, statement_text, size=54, bold=True, color='FFFFFF',
         align='c', anchor='m', ls=1.1)
    return sl


def recap(prs, rows, title='Every Git tool fixes one problem you felt',
          heads=('We felt…', 'We invented…', 'Git calls it')):
    """rows: list of (pain, invention, git_cmd). <=9 rows."""
    sl = blank(prs)
    text(sl, M, 0.45, CW, 0.85, title, size=40, bold=True, anchor='m')
    cols = [(M, 4.55), (M + 4.75, 4.0), (M + 8.95, 2.88)]
    y0, rh = 1.55, 0.54
    for (cx, cw), hd in zip(cols, heads):
        text(sl, cx, y0, cw, 0.4, hd.upper(), size=15, bold=True, color=C['MUTED'], anchor='m')
    for i, (a, b, c) in enumerate(rows):
        yy = y0 + 0.45 + i * rh
        arrow(sl, M, yy, W - M, yy, color=C['HAIR'], w=1, head=False)
        text(sl, cols[0][0], yy, cols[0][1], rh, a, size=20, anchor='m')
        text(sl, cols[1][0], yy, cols[1][1], rh, b, size=20, color=C['INK2'], anchor='m')
        text(sl, cols[2][0], yy, cols[2][1], rh, c, size=20, bold=True, color=C['PURPLE'],
             font=MONO, anchor='m')
    return sl


def pause(prs, question, rhythm='Think 1 min  ·  Pair 2 min  ·  Share'):
    sl = blank(prs)
    text(sl, M, 1.0, CW, 1.0, '💬', size=60, align='c', anchor='m')
    text(sl, 1.4, 2.2, W - 2.8, 2.9, question, size=44, bold=True, align='c', anchor='m',
         ls=1.08)
    chip(sl, (W - 6.2) / 2, 5.5, 6.2, 0.56, rhythm, fill=C['LAV'], color=C['INK'], size=20)
    return sl


def break_slide(prs, minutes=5, back_at='__:__'):
    sl = blank(prs, C['YELLOW'])
    text(sl, M, 2.0, CW, 1.1, '☕', size=72, align='c', anchor='m')
    text(sl, M, 3.15, CW, 1.2, f'{minutes}-minute break', size=60, bold=True, align='c',
         anchor='m')
    text(sl, M, 4.45, CW, 0.7, f'Back at {back_at}', size=28, color=C['INK'], align='c')
    return sl
